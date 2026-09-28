#!/usr/bin/env python3
"""Typeset the hand-written materials of the vendor-quote-evaluation task.

The quotation letters and due-diligence notes below were written by hand.
This script only puts that prose into the two Word documents the agent
receives (``initial_workspace/Vendor_Quotations.docx`` and
``initial_workspace/Supplier_Due_Diligence_Notes.docx``) so the materials can
be edited in one place and regenerated consistently.  It never computes
anything: the ground truth is derived from the finished documents by
``derive_groundtruth.py``.

    python author_materials.py            # rewrites the two .docx files
"""

import os

from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TASK_DIR = os.path.dirname(HERE)
INITIAL_WORKSPACE = os.path.join(TASK_DIR, "initial_workspace")

QUOTATIONS_FILE = "Vendor_Quotations.docx"
NOTES_FILE = "Supplier_Due_Diligence_Notes.docx"

# --------------------------------------------------------------------------- #
# Quotations, in the order they were received (one entry per quotation letter).
# Ashford Seating & Furniture sent two revisions: entry 2 (superseded) and
# entry 10 (current).
# --------------------------------------------------------------------------- #
QUOTATIONS = [
    """Vendor Code: NPS
Vendor: Northpeak Office Supply
Quotation Reference: NPS-Q-4471
Quotation Date: 2 September 2025
Item Offered: Northpeak Aria mesh-back task chair, model AR-200, with adjustable lumbar support and 4D armrests
Quantity Quoted: 150 units
Unit Price: USD 318.00 per chair
Volume Discount: 5% on goods for orders of 100 units or more
Freight: USD 1,200.00 for a single delivery to the Riverside site
Lead Time: 30 calendar days after receipt of purchase order
Warranty: 5 years on mechanism, 2 years on fabric and foam
Certifications: BIFMA X5.1 certified; GREENGUARD Gold
Quote Validity: Valid until 30 November 2025
Payment Terms: Net 30 days from delivery
Notes: Delivery includes unpacking and removal of packaging. Workstation assembly is available at USD 12.00 per chair on request and is not included in the prices above.""",

    """Vendor Code: ASF
Vendor: Ashford Seating & Furniture
Quotation Reference: ASF-2509-01
Quotation Date: 5 September 2025
Item Offered: Ashford Contour Pro task chair with synchro-tilt mechanism and adjustable seat depth
Quantity Quoted: 150 units
Unit Price: USD 305.00 per chair
Volume Discount: None
Freight: USD 1,500.00 to the Riverside site
Lead Time: 60 calendar days after receipt of purchase order
Warranty: 6 years on mechanism and base
Certifications: BIFMA X5.1 certified
Quote Validity: Valid until 31 December 2025
Payment Terms: Net 45 days from invoice
Notes: The lead time reflects our current factory backlog; we will advise you if capacity frees up earlier.""",

    """Vendor Code: BVL
Vendor: Blue Valley Logistics & Furnishing
Quotation Reference: BVL/0906/RIV
Quotation Date: 6 September 2025
Item Offered: BV FlexBack task chair, mesh back, polished aluminium base
Quantity Quoted: 150 units
Unit Price: USD 279.00 per chair
Volume Discount: None (prices quoted are already net)
Freight: USD 2,400.00 in two shipments to the Riverside site
Lead Time: 50 calendar days after receipt of purchase order
Warranty: 3 years on mechanism, 1 year on fabric
Certifications: BIFMA X5.1 certified
Quote Validity: Valid until 1 December 2025
Payment Terms: 50% deposit with order, balance on delivery
Notes: Two shipments are required because of container availability; both are included in the freight figure above.""",

    """Vendor Code: SOL
Vendor: Solstice Contract Interiors
Quotation Reference: SCI-Q-2025-118
Quotation Date: 8 September 2025
Item Offered: Solstice Meridian task chair, mesh back, height-adjustable arms
Quantity Quoted: 150 units
Unit Price: USD 289.60 per chair
Volume Discount: 5% on goods for orders of 100 units or more
Freight: USD 1,100.00 for delivery to the Riverside site
Lead Time: 38 calendar days after receipt of purchase order
Warranty: 5 years on mechanism
Certifications: BIFMA X5.1 certified
Quote Validity: Valid for 30 days from the quotation date
Payment Terms: Net 30 days from invoice
Notes: Pricing is based on our September production slot; an extension of the validity period can be requested in writing.""",

    """Vendor Code: HMP
Vendor: Harbor & Main Procurement
Quotation Reference: HM-25-0912
Quotation Date: 9 September 2025
Item Offered: Harbor Line Ergo 7 task chair with adjustable lumbar support and seat slide
Quantity Quoted: 150 units
Unit Price: USD 299.00 per chair
Volume Discount: 10% on goods for orders of 200 units or more
Freight: USD 1,750.00 to the Riverside site
Lead Time: 42 calendar days after receipt of purchase order
Warranty: 60 months on mechanism and gas lift
Certifications: BIFMA X5.1 certified
Quote Validity: Valid until 31 October 2025
Payment Terms: Net 30 days from invoice
Notes: Should the order be increased to 200 chairs, the 10% volume discount would apply to the whole order.""",

    """Vendor Code: KWE
Vendor: Kestrel Workspace Europe GmbH
Quotation Reference: KWE-EX-3308
Quotation Date: 10 September 2025
Item Offered: Kestrel Motion S2 task chair, mesh back, EN 1335 tested
Quantity Quoted: 150 units
Unit Price: EUR 282.50 per unit
Volume Discount: None
Freight: EUR 900.00, sea freight delivered to the Riverside site (DAP)
Lead Time: 35 calendar days after receipt of purchase order
Warranty: 5 years on mechanism
Certifications: BIFMA X5.1 certified; EN 1335
Quote Validity: Valid until 15 November 2025
Payment Terms: Net 30 days from invoice
Notes: All prices are in euro. Import duties, if any, are for the buyer's account.""",

    """Vendor Code: GRD
Vendor: Granite Ridge Distribution
Quotation Reference: GRD-Q-7720
Quotation Date: 12 September 2025
Item Offered: Granite Ridge Summit task chair, mesh back, adjustable arms
Quantity Quoted: 150 units
Unit Price: USD 296.00 per chair
Volume Discount: None
Freight: Included in the unit price (delivered to the Riverside site)
Lead Time: 45 calendar days after receipt of purchase order
Warranty: 60 months on mechanism
Certifications: BIFMA X5.1 certified
Quote Validity: Valid until 30 November 2025
Payment Terms: Net 30 days from invoice
Notes: Our quality management system is currently being re-certified; documentation can be provided on request.""",

    """Vendor Code: TRV
Vendor: Trevane Commercial Furniture
Quotation Reference: TCF-Q-0915
Quotation Date: 15 September 2025
Item Offered: Trevane Axis task chair with mesh back and adjustable headrest
Quantity Quoted: 120 units (the balance of 30 units can be supplied from Q1 2026)
Unit Price: USD 301.00 per chair
Volume Discount: 3% on goods for orders of 100 units or more
Freight: USD 950.00 to the Riverside site
Lead Time: 28 calendar days after receipt of purchase order
Warranty: 7 years on all components
Certifications: BIFMA X5.1 certified
Quote Validity: Valid until 31 December 2025
Payment Terms: Net 30 days from invoice
Notes: 120 chairs are available from current stock; the remaining 30 would follow in the first quarter of 2026 at the same unit price.""",

    """Vendor Code: PRM
Vendor: Premier Ergonomics Inc.
Quotation Reference: PE-25-2209
Quotation Date: 17 September 2025
Item Offered: Premier Apex executive task chair, mesh back, adjustable headrest and 4D arms
Quantity Quoted: 150 units
Unit Price: USD 412.00 per chair
Volume Discount: 5% on goods for orders of 100 units or more
Freight: USD 1,900.00 to the Riverside site
Lead Time: 20 calendar days after receipt of purchase order
Warranty: 10 years on mechanism, 5 years on fabric
Certifications: BIFMA X5.1 certified; GREENGUARD Gold
Quote Validity: Valid until 31 December 2025
Payment Terms: Net 30 days from invoice
Notes: Price includes white-glove delivery and workstation assembly.""",

    """Vendor Code: ASF
Vendor: Ashford Seating & Furniture
Quotation Reference: ASF-2509-01-R2
Quotation Date: 19 September 2025
Item Offered: Ashford Contour Pro task chair with synchro-tilt mechanism and adjustable seat depth
Quantity Quoted: 150 units
Unit Price: USD 309.50 per chair
Volume Discount: 4% on goods for orders of 150 units or more
Freight: Included in the unit price (single delivery to the Riverside site)
Lead Time: 40 calendar days after receipt of purchase order
Warranty: 6 years on mechanism and base
Certifications: BIFMA X5.1 certified
Quote Validity: Valid until 31 December 2025
Payment Terms: Net 45 days from invoice
Notes: This quotation supersedes our quotation ASF-2509-01 dated 5 September 2025 in its entirety.""",
]

# --------------------------------------------------------------------------- #
# Due-diligence notes, alphabetical by registered name (one entry per vendor).
# --------------------------------------------------------------------------- #
NOTES = [
    """Vendor Code: ASF
Registered Name: Ashford Seating & Furniture Ltd.
ISO 9001: Certificate AS-9001-0871, valid until 9 February 2027
Reference Checks: Two references contacted (Harlow Insurance, Pinecrest College); both satisfactory
Delivery History: One previous order (2022 training centre) delivered one week late because of port congestion
Financial Standing: Credit rating A-; no adverse findings
Notes: The vendor confirmed that revised pricing was issued after our clarification call on 18 September 2025.""",

    """Vendor Code: BVL
Registered Name: Blue Valley Logistics & Furnishing LLC
ISO 9001: Certificate BV-22-3190, valid until 30 April 2026
Reference Checks: Two references contacted; one reported a late delivery in 2024
Delivery History: No previous orders with us
Financial Standing: Credit rating B+; no adverse findings
Notes: Consistently the lowest pricing in previous tenders; warranty terms have been raised as a concern before.""",

    """Vendor Code: GRD
Registered Name: Granite Ridge Distribution Co.
ISO 9001: Certificate GR-2019-0455 expired on 30 June 2024; the vendor states that a renewal audit is planned but no date has been confirmed
Reference Checks: Two references contacted; both satisfactory
Delivery History: One previous order (2021 call centre) delivered on schedule
Financial Standing: Credit rating A-; no adverse findings
Notes: Freight is normally included in Granite Ridge pricing.""",

    """Vendor Code: HMP
Registered Name: Harbor & Main Procurement Inc.
ISO 9001: Certificate HM-9001-1160, valid until 20 October 2025 (renewal audit completed, new certificate pending)
Reference Checks: Two references contacted; both satisfactory
Delivery History: No previous orders with us
Financial Standing: Credit rating A; no adverse findings
Notes: None.""",

    """Vendor Code: KWE
Registered Name: Kestrel Workspace Europe GmbH
ISO 9001: Certificate DE-9001-55821, valid until 3 August 2027
Reference Checks: Two references contacted; both satisfactory
Delivery History: No previous orders with us; the vendor quoted in euro in earlier tenders as well
Financial Standing: Credit rating A; no adverse findings
Notes: No import duties are expected for this product category.""",

    """Vendor Code: NPS
Registered Name: Northpeak Office Supply Ltd.
ISO 9001: Certificate NP-9001-2211, valid until 14 March 2027
Reference Checks: Two references contacted (Meridian Bank, Calder Logistics); both satisfactory
Delivery History: Two previous orders, both delivered on schedule
Financial Standing: Credit rating A; no adverse findings
Notes: Preferred supplier for the 2023 headquarters refurbishment.""",

    """Vendor Code: PRM
Registered Name: Premier Ergonomics Inc.
ISO 9001: Certificate PE-0932, valid until 11 January 2028
Reference Checks: Two references contacted; both satisfactory
Delivery History: No previous orders with us
Financial Standing: Credit rating A; no adverse findings
Notes: Positioned as a premium brand; pricing has historically exceeded our budgets.""",

    """Vendor Code: SOL
Registered Name: Solstice Contract Interiors
ISO 9001: Certificate SC-9001-7714, valid until 27 May 2026
Reference Checks: Two references contacted; both satisfactory
Delivery History: One previous order (2024 meeting rooms) delivered on schedule
Financial Standing: Credit rating A-; no adverse findings
Notes: The sales representative mentioned that quotation validity periods are kept short because of fabric price volatility.""",

    """Vendor Code: TRV
Registered Name: Trevane Commercial Furniture
ISO 9001: Certificate TC-9001-4402, valid until 18 September 2026
Reference Checks: Two references contacted; both satisfactory
Delivery History: No previous orders with us
Financial Standing: Credit rating B+; no adverse findings
Notes: Stock levels are limited; the vendor usually delivers in tranches.""",
]


def _new_document():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    return doc


def _add_entries(doc, label, entries):
    for index, block in enumerate(entries, start=1):
        marker = doc.add_paragraph()
        marker.add_run(f"{label} {index}").bold = True
        doc.add_paragraph(block)  # "\n" becomes a line break inside one paragraph
        doc.add_paragraph("")


def write_quotations(path):
    doc = _new_document()
    doc.add_heading("Quotations received for RFQ-2025-031", level=1)
    doc.add_paragraph(
        "Ergonomic task chairs for the Riverside office expansion. Compiled by "
        "Procurement on 22 September 2025; entries are listed in the order in "
        "which they were received."
    )
    _add_entries(doc, "Quotation", QUOTATIONS)
    doc.save(path)


def write_notes(path):
    doc = _new_document()
    doc.add_heading("Supplier due-diligence notes for RFQ-2025-031", level=1)
    doc.add_paragraph(
        "Prepared by Procurement on 20 September 2025 from the supplier "
        "registration files, certificate copies and reference calls. Suppliers "
        "are listed alphabetically by registered name."
    )
    _add_entries(doc, "Supplier", NOTES)
    doc.save(path)


def main():
    os.makedirs(INITIAL_WORKSPACE, exist_ok=True)
    quotations_path = os.path.join(INITIAL_WORKSPACE, QUOTATIONS_FILE)
    notes_path = os.path.join(INITIAL_WORKSPACE, NOTES_FILE)
    write_quotations(quotations_path)
    write_notes(notes_path)
    print(f"wrote {quotations_path} ({len(QUOTATIONS)} quotations)")
    print(f"wrote {notes_path} ({len(NOTES)} suppliers)")


if __name__ == "__main__":
    main()
