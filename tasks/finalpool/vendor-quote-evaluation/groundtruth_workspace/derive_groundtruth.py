#!/usr/bin/env python3
"""Derive the ground truth of the vendor-quote-evaluation task from its materials.

Reads only what the agent sees in ``initial_workspace/``:

* ``Evaluation_Rules.md``               - requirement constants and decision rules
* ``Vendor_Quotations.docx``            - the hand-written quotation letters
* ``Supplier_Due_Diligence_Notes.docx`` - the hand-written due-diligence notes

and writes into this directory:

* ``Quotation_Evaluation_<CODE>.docx``  - one expected evaluation sheet per vendor
                                          (latest revision only), laid out as
                                          ``Format.md`` prescribes
* ``award.txt``                         - name of the vendor to be awarded
* ``expected_results.json``             - the full derivation (parsed fields,
                                          totals, criteria, ranking) for the
                                          README and the tests

Usage:
    python derive_groundtruth.py            # regenerate everything
    python derive_groundtruth.py --print    # only print the derivation table

``check_trap_sensitivity.py`` imports ``evaluate()`` and re-runs it with
deliberate mistakes switched on to prove that every planted trap changes the
result.  Money is handled with ``Decimal``; all figures in the materials were
chosen so that every intermediate value is an exact number of cents.
"""

from __future__ import annotations

import argparse
import glob
import json
import os
import re
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from typing import Dict, FrozenSet, List, Optional, Tuple

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TASK_DIR = os.path.dirname(HERE)
INITIAL_WORKSPACE = os.path.join(TASK_DIR, "initial_workspace")

RULES_FILE = "Evaluation_Rules.md"
QUOTATIONS_FILE = "Vendor_Quotations.docx"
NOTES_FILE = "Supplier_Due_Diligence_Notes.docx"

OUTPUT_PREFIX = "Quotation_Evaluation_"
AWARD_FILE = "award.txt"
RESULTS_FILE = "expected_results.json"

HEADING_TEXT = "Quotation Evaluation Sheet"
TABLE_HEADER = [
    "Total Evaluated Cost (USD)",
    "Mandatory Criteria Met",
    "Decision",
    "Disqualification Reasons",
]
DECISION_AWARD = "Award"
DECISION_NOT_SELECTED = "Compliant - Not Selected"
DECISION_DISQUALIFIED = "Disqualified"
NO_REASONS = "None"

# Reason phrases in criterion order; they must appear verbatim in the rules.
CRITERIA = [
    "Partial quantity",
    "Lead time exceeds 45 days",
    "Warranty below 5 years",
    "Missing or expired certification",
    "Quote expires before award date",
    "Exceeds budget ceiling",
]

CENT = Decimal("0.01")

# Deliberate mistakes an agent may make; each must change the derived result.
MISTAKES = {
    "use_first_revision": "evaluate a vendor's earliest quotation instead of the latest",
    "ignore_currency": "treat EUR figures as if they were USD",
    "ignore_discount_threshold": "apply every volume discount regardless of its quantity threshold",
    "skip_discounts": "never apply a volume discount",
    "discount_on_freight": "apply the volume discount to freight as well",
    "ignore_due_diligence": "assume every vendor's ISO 9001 certificate is valid",
    "ignore_validity_duration": "treat a validity given in days as always valid",
    "ignore_warranty": "skip the warranty criterion",
    "ignore_quantity": "skip the quantity criterion",
    "ignore_lead_time": "skip the lead-time criterion",
    "ignore_budget": "skip the budget criterion",
}


# --------------------------------------------------------------------------- #
# parsing helpers
# --------------------------------------------------------------------------- #
def parse_day(text: str) -> date:
    return datetime.strptime(text.strip(), "%d %B %Y").date()


def _field(block: str, label: str) -> str:
    match = re.search(rf"^{re.escape(label)}:\s*(.+?)\s*$", block, re.M)
    if not match:
        raise ValueError(f"field '{label}' not found in block starting {block[:40]!r}")
    return match.group(1)


def _money(text: str) -> Tuple[str, Decimal]:
    match = re.search(r"\b(USD|EUR)\s*([\d,]+\.\d{2})", text)
    if not match:
        raise ValueError(f"no amount in {text!r}")
    return match.group(1), Decimal(match.group(2).replace(",", ""))


def parse_rules(text: str) -> Dict:
    def grab(pattern: str) -> str:
        match = re.search(pattern, text)
        if not match:
            raise ValueError(f"rule not found: {pattern}")
        return match.group(1)

    rules = {
        "quantity_required": int(grab(r"Quantity required:\s*(\d+)\s*units")),
        "award_date": parse_day(grab(r"Planned award date:\s*(\d{1,2} [A-Z][a-z]+ \d{4})")),
        "budget_ceiling": Decimal(grab(r"Budget ceiling:\s*USD\s*([\d,]+\.\d{2})").replace(",", "")),
        "max_lead_days": int(grab(r"Maximum lead time:\s*(\d+)\s*calendar days")),
        "min_warranty_months": 12 * int(grab(r"Minimum warranty on the chair mechanism:\s*(\d+)\s*years")),
        "eur_usd": Decimal(grab(r"EUR 1\.00 = USD\s*(\d+\.\d+)")),
    }
    for phrase in CRITERIA:
        if phrase not in text:
            raise ValueError(f"reason phrase {phrase!r} missing from {RULES_FILE}")
    return rules


def extract_entries(docx_path: str, label: str) -> List[str]:
    """Return the text block that follows every '<label> N' marker paragraph."""
    paragraphs = [p.text for p in Document(docx_path).paragraphs]
    blocks = []
    i = 0
    while i < len(paragraphs):
        if re.fullmatch(rf"{label} \d+", paragraphs[i].strip()):
            j = i + 1
            while j < len(paragraphs) and not paragraphs[j].strip():
                j += 1
            blocks.append(paragraphs[j].strip())
            i = j
        i += 1
    return blocks


def parse_quotation(block: str) -> Dict:
    q: Dict = {"block": block}
    q["code"] = _field(block, "Vendor Code")
    q["vendor"] = _field(block, "Vendor")
    q["reference"] = _field(block, "Quotation Reference")
    q["date"] = parse_day(_field(block, "Quotation Date"))
    q["quantity"] = int(re.match(r"(\d+)\s*units", _field(block, "Quantity Quoted")).group(1))
    q["currency"], q["unit_price"] = _money(_field(block, "Unit Price"))

    discount = _field(block, "Volume Discount")
    match = re.match(r"(\d+(?:\.\d+)?)%\s+on goods for orders of\s+(\d+)\s+units or more", discount)
    q["discount_pct"] = Decimal(match.group(1)) if match else None
    q["discount_threshold"] = int(match.group(2)) if match else None
    if not match and not discount.lower().startswith("none"):
        raise ValueError(f"unrecognised discount clause: {discount!r}")

    freight = _field(block, "Freight")
    if freight.lower().startswith("included"):
        q["freight_included"] = True
        q["freight"] = Decimal("0.00")
    else:
        q["freight_included"] = False
        currency, amount = _money(freight)
        if currency != q["currency"]:
            raise ValueError(f"freight currency differs from unit price in {q['reference']}")
        q["freight"] = amount

    q["lead_days"] = int(re.match(r"(\d+)\s*calendar days", _field(block, "Lead Time")).group(1))

    warranty = _field(block, "Warranty")
    match = (re.search(r"(\d+)\s*(years?|months?)\s+on\s+mechanism", warranty)
             or re.search(r"(\d+)\s*(years?|months?)", warranty))
    number = int(match.group(1))
    q["warranty_months"] = number * 12 if match.group(2).startswith("year") else number

    q["bifma"] = "BIFMA X5.1" in _field(block, "Certifications")

    validity = _field(block, "Quote Validity")
    match = re.match(r"Valid until (\d{1,2} [A-Z][a-z]+ \d{4})", validity)
    if match:
        q["validity_kind"] = "date"
        q["valid_until"] = parse_day(match.group(1))
    else:
        match = re.match(r"Valid for (\d+) days from the quotation date", validity)
        if not match:
            raise ValueError(f"unrecognised validity clause: {validity!r}")
        q["validity_kind"] = "duration"
        q["valid_until"] = q["date"] + timedelta(days=int(match.group(1)))
    return q


def parse_notes(block: str) -> Dict:
    n: Dict = {"block": block}
    n["code"] = _field(block, "Vendor Code")
    n["registered_name"] = _field(block, "Registered Name")
    iso = _field(block, "ISO 9001")
    match = re.search(r"valid until (\d{1,2} [A-Z][a-z]+ \d{4})", iso)
    n["iso_valid_until"] = parse_day(match.group(1)) if match else None
    match = re.search(r"expired on (\d{1,2} [A-Z][a-z]+ \d{4})", iso)
    n["iso_expired_on"] = parse_day(match.group(1)) if match else None
    return n


def load_materials(initial_workspace: str = INITIAL_WORKSPACE):
    with open(os.path.join(initial_workspace, RULES_FILE), encoding="utf-8") as fh:
        rules = parse_rules(fh.read())
    quotations = [parse_quotation(b) for b in extract_entries(os.path.join(initial_workspace, QUOTATIONS_FILE), "Quotation")]
    notes = {}
    for block in extract_entries(os.path.join(initial_workspace, NOTES_FILE), "Supplier"):
        note = parse_notes(block)
        if note["code"] in notes:
            raise ValueError(f"duplicate due-diligence entry for {note['code']}")
        notes[note["code"]] = note
    missing = {q["code"] for q in quotations} - set(notes)
    if missing:
        raise ValueError(f"no due-diligence entry for {sorted(missing)}")
    return rules, quotations, notes


# --------------------------------------------------------------------------- #
# evaluation
# --------------------------------------------------------------------------- #
def evaluate(rules: Dict, quotations: List[Dict], notes: Dict[str, Dict],
             mistakes: FrozenSet[str] = frozenset()) -> Tuple[List[Dict], Optional[Dict]]:
    """Apply Evaluation_Rules.md; ``mistakes`` switches deliberate errors on."""
    unknown = set(mistakes) - set(MISTAKES)
    if unknown:
        raise ValueError(f"unknown mistakes {sorted(unknown)}")

    revisions: Dict[str, List[Dict]] = {}
    for q in quotations:
        revisions.setdefault(q["code"], []).append(q)

    results: List[Dict] = []
    for code, revs in revisions.items():
        revs = sorted(revs, key=lambda item: item["date"])
        q = revs[0] if "use_first_revision" in mistakes else revs[-1]
        note = notes[code]

        goods = q["unit_price"] * q["quantity"]
        discount_applies = False
        if q["discount_pct"] is not None:
            discount_applies = q["quantity"] >= q["discount_threshold"]
            if "ignore_discount_threshold" in mistakes:
                discount_applies = True
            if "skip_discounts" in mistakes:
                discount_applies = False
        rate = q["discount_pct"] / Decimal(100) if discount_applies else Decimal(0)
        discount = goods * rate
        freight = q["freight"]
        total_native = goods - discount + freight
        if "discount_on_freight" in mistakes:
            total_native = (goods + freight) * (Decimal(1) - rate)
        total_usd = total_native
        if q["currency"] == "EUR" and "ignore_currency" not in mistakes:
            total_usd = total_native * rules["eur_usd"]
        total_usd = total_usd.quantize(CENT, rounding=ROUND_HALF_UP)

        iso_ok = note["iso_valid_until"] is not None and note["iso_valid_until"] >= rules["award_date"]
        if "ignore_due_diligence" in mistakes:
            iso_ok = True
        validity_ok = q["valid_until"] >= rules["award_date"]
        if "ignore_validity_duration" in mistakes and q["validity_kind"] == "duration":
            validity_ok = True

        checks = [
            q["quantity"] >= rules["quantity_required"] or "ignore_quantity" in mistakes,
            q["lead_days"] <= rules["max_lead_days"] or "ignore_lead_time" in mistakes,
            q["warranty_months"] >= rules["min_warranty_months"] or "ignore_warranty" in mistakes,
            iso_ok and q["bifma"],
            validity_ok,
            total_usd <= rules["budget_ceiling"] or "ignore_budget" in mistakes,
        ]
        reasons = [phrase for phrase, ok in zip(CRITERIA, checks) if not ok]

        results.append({
            "code": code,
            "vendor": q["vendor"],
            "registered_name": note["registered_name"],
            "quotation_reference": q["reference"],
            "quotation_date": q["date"].isoformat(),
            "superseded_references": [r["reference"] for r in revs if r is not q],
            "quantity_quoted": q["quantity"],
            "currency": q["currency"],
            "unit_price": str(q["unit_price"]),
            "discount_pct": str(q["discount_pct"]) if q["discount_pct"] is not None else None,
            "discount_threshold": q["discount_threshold"],
            "discount_applied": discount_applies,
            "goods_subtotal": str(goods.quantize(CENT)),
            "discount_amount": str(discount.quantize(CENT)),
            "freight": str(freight),
            "freight_included": q["freight_included"],
            "total_native": str(total_native.quantize(CENT)),
            "total_usd": str(total_usd),
            "lead_days": q["lead_days"],
            "warranty_months": q["warranty_months"],
            "bifma": q["bifma"],
            "iso_valid_until": note["iso_valid_until"].isoformat() if note["iso_valid_until"] else None,
            "iso_ok": iso_ok,
            "validity_kind": q["validity_kind"],
            "valid_until": q["valid_until"].isoformat(),
            "criteria": [{"reason_if_failed": phrase, "met": ok} for phrase, ok in zip(CRITERIA, checks)],
            "criteria_met": sum(1 for ok in checks if ok),
            "reasons": reasons,
            "quotation_block": q["block"],
            "notes_block": note["block"],
            "_total": total_usd,
        })

    compliant = sorted((r for r in results if not r["reasons"]),
                       key=lambda r: (r["_total"], r["lead_days"], r["vendor"]))
    for rank, r in enumerate(compliant, start=1):
        r["rank"] = rank
    for r in results:
        r.setdefault("rank", None)
        if r["reasons"]:
            r["decision"] = DECISION_DISQUALIFIED
        elif r["rank"] == 1:
            r["decision"] = DECISION_AWARD
        else:
            r["decision"] = DECISION_NOT_SELECTED
        del r["_total"]
    results.sort(key=lambda r: r["code"])
    award = compliant[0] if compliant else None
    return results, award


def format_usd(amount: str | Decimal) -> str:
    return f"{Decimal(amount):,.2f}"


def table_row(result: Dict) -> List[str]:
    return [
        format_usd(result["total_usd"]),
        f"{result['criteria_met']}/{len(CRITERIA)}",
        result["decision"],
        "; ".join(result["reasons"]) if result["reasons"] else NO_REASONS,
    ]


# --------------------------------------------------------------------------- #
# writing the expected documents
# --------------------------------------------------------------------------- #
def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _fill_cell(cell, text: str, fill: str) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _shade(cell, fill)


def _title(doc, text: str) -> None:
    doc.add_paragraph().add_run(text).bold = True


def write_sheet(path: str, result: Dict) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Cambria"
    normal.font.size = Pt(11)

    heading = doc.add_heading(HEADING_TEXT, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in heading.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(14)
        run.font.bold = True

    doc.add_paragraph(f"Vendor Code: {result['code']}")
    doc.add_paragraph(f"Vendor Name: {result['vendor']}")
    _title(doc, "Quotation Summary:")
    doc.add_paragraph(result["quotation_block"])
    _title(doc, "Due Diligence Notes:")
    doc.add_paragraph(result["notes_block"])
    _title(doc, "Evaluation Conclusion:")

    table = doc.add_table(rows=2, cols=len(TABLE_HEADER))
    table.style = "Table Grid"
    for j, text in enumerate(TABLE_HEADER):
        _fill_cell(table.cell(0, j), text, "D9D9D9")
    for j, text in enumerate(table_row(result)):
        _fill_cell(table.cell(1, j), text, "DAEEF3")
    doc.save(path)


def print_table(results: List[Dict], award: Optional[Dict]) -> None:
    print(f"{'code':<5}{'vendor':<36}{'qty':>4}{'total USD':>12}{'lead':>6}{'warr':>6}  met  decision")
    for r in results:
        print(f"{r['code']:<5}{r['vendor']:<36}{r['quantity_quoted']:>4}{format_usd(r['total_usd']):>12}"
              f"{r['lead_days']:>6}{r['warranty_months']:>6}  {r['criteria_met']}/6  {r['decision']}"
              + (f"  [{'; '.join(r['reasons'])}]" if r["reasons"] else ""))
    print("award:", award["vendor"] if award else "none")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--print", action="store_true", help="print the derivation and exit")
    args = parser.parse_args()

    rules, quotations, notes = load_materials()
    results, award = evaluate(rules, quotations, notes)
    print_table(results, award)
    if args.print:
        return
    if award is None:
        raise SystemExit("no compliant vendor; the materials must define a winner")

    for stale in glob.glob(os.path.join(HERE, f"{OUTPUT_PREFIX}*.docx")):
        os.remove(stale)
    for r in results:
        write_sheet(os.path.join(HERE, f"{OUTPUT_PREFIX}{r['code']}.docx"), r)
    with open(os.path.join(HERE, AWARD_FILE), "w", encoding="utf-8", newline="\n") as fh:
        fh.write(award["vendor"] + "\n")

    payload = {
        "rules": {k: (v.isoformat() if isinstance(v, date) else str(v)) for k, v in rules.items()},
        "criteria_order": CRITERIA,
        "award": {"code": award["code"], "vendor": award["vendor"], "total_usd": award["total_usd"]},
        "ranking_of_compliant_vendors": [r["code"] for r in sorted(
            (r for r in results if r["rank"]), key=lambda r: r["rank"])],
        "vendors": results,
    }
    with open(os.path.join(HERE, RESULTS_FILE), "w", encoding="utf-8", newline="\n") as fh:
        json.dump(payload, fh, indent=2, ensure_ascii=False)
    print(f"wrote {len(results)} sheets, {AWARD_FILE} and {RESULTS_FILE} into {HERE}")


if __name__ == "__main__":
    main()
