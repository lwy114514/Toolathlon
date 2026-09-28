#!/usr/bin/env python3
"""Self-test for evaluation/check_local.py.

Builds throw-away agent workspaces from the ground truth, mutates them in the
ways an agent tends to get wrong (or merely formats differently), and asserts
that the checker accepts exactly the correct variants.

Run from the repository root or the task directory:
    python tasks/finalpool/vendor-quote-evaluation/evaluation/test_check_local.py
"""

import contextlib
import io
import json
import os
import re
import shutil
import sys
import tempfile

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from check_local import AWARD_FILE, OUTPUT_PREFIX, check_local  # noqa: E402

TASK_DIR = os.path.dirname(HERE)
GT_DIR = os.path.join(TASK_DIR, "groundtruth_workspace")
INITIAL_DIR = os.path.join(TASK_DIR, "initial_workspace")

with open(os.path.join(GT_DIR, "expected_results.json"), encoding="utf-8") as fh:
    EXPECTED = json.load(fh)
VENDORS = {v["code"]: v for v in EXPECTED["vendors"]}
AWARD = EXPECTED["award"]
CODES = sorted(VENDORS)

# paragraph indices inside a sheet (heading, code, name, title, block, title, block, title)
IDX_VENDOR_NAME = 2
IDX_QUOTATION_BLOCK = 4

results = []


def case(name, expect_pass, builder):
    tmp = tempfile.mkdtemp(prefix="vqe_")
    try:
        builder(tmp)
        with contextlib.redirect_stdout(io.StringIO()):
            ok, msg = check_local(tmp, GT_DIR)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    status = "ok " if ok == expect_pass else "BAD"
    results.append(ok == expect_pass)
    print(f"[{status}] {name:<58} -> {'pass' if ok else 'fail'}: {msg[:150]}")


# ------------------------------------------------------------------ builders
def sheet_name(code):
    return f"{OUTPUT_PREFIX}{code}.docx"


def copy_gt(tmp):
    for code in CODES:
        shutil.copy(os.path.join(GT_DIR, sheet_name(code)), os.path.join(tmp, sheet_name(code)))
    shutil.copy(os.path.join(GT_DIR, AWARD_FILE), os.path.join(tmp, AWARD_FILE))


def set_award(tmp, text):
    with open(os.path.join(tmp, AWARD_FILE), "w", encoding="utf-8") as fh:
        fh.write(text)


def gt_parts(code):
    doc = Document(os.path.join(GT_DIR, sheet_name(code)))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    rows = [[cell.text for cell in row.cells] for row in doc.tables[0].rows]
    return paragraphs, rows


def write_sheet(tmp, code, paragraphs, rows, split_lines=False):
    """Write a plain sheet (no styles, no shading) from paragraph texts and table rows."""
    doc = Document()
    for text in paragraphs:
        if split_lines:
            for line in text.split("\n"):
                doc.add_paragraph(line)
        else:
            doc.add_paragraph(text)
    if rows is not None:
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for i, row in enumerate(rows):
            for j, cell in enumerate(row):
                table.cell(i, j).text = cell
    doc.save(os.path.join(tmp, sheet_name(code)))


def modified(code, paragraphs=None, rows=None, split_lines=False):
    """Builder: ground truth with one sheet rewritten through the given transforms."""
    def build(tmp):
        copy_gt(tmp)
        paras, table = gt_parts(code)
        if paragraphs is not None:
            paras = paragraphs(list(paras))
        if rows is not None:
            table = rows([list(r) for r in table])
        write_sheet(tmp, code, paras, table, split_lines=split_lines)
    return build


def with_data_row(**changes):
    """Table transform that overwrites data-row cells: cost, met, decision, reasons."""
    position = {"cost": 0, "met": 1, "decision": 2, "reasons": 3}

    def transform(rows):
        for key, value in changes.items():
            rows[1][position[key]] = value
        return rows
    return transform


def quotation_block(index):
    """The block that follows 'Quotation <index>' in the materials (2 = superseded ASF revision)."""
    paragraphs = [p.text for p in Document(os.path.join(INITIAL_DIR, "Vendor_Quotations.docx")).paragraphs]
    for i, text in enumerate(paragraphs):
        if text.strip() == f"Quotation {index}":
            j = i + 1
            while not paragraphs[j].strip():
                j += 1
            return paragraphs[j].strip()
    raise AssertionError(f"Quotation {index} not found")


def messy_but_correct(tmp):
    """Every sheet rewritten without formatting, one paragraph per line, other spellings."""
    for code in CODES:
        paras, rows = gt_parts(code)
        cost, met, decision, reasons = rows[1]
        rows[1] = [
            f"USD {cost}",
            met.replace("/", " of "),
            decision.lower(),
            ", ".join(reversed(reasons.split("; "))),
        ]
        write_sheet(tmp, code, paras, rows, split_lines=True)
    set_award(tmp, AWARD["code"])


def registered_name_award(tmp):
    copy_gt(tmp)
    set_award(tmp, VENDORS[AWARD["code"]]["registered_name"] + "\n\n")


def name_and_code_award(tmp):
    copy_gt(tmp)
    set_award(tmp, f"{AWARD['vendor']} ({AWARD['code']})\n")


def extra_files(tmp):
    copy_gt(tmp)
    with open(os.path.join(tmp, "working_notes.txt"), "w", encoding="utf-8") as fh:
        fh.write("scratch\n")
    shutil.copy(os.path.join(GT_DIR, sheet_name(CODES[0])), os.path.join(tmp, f"{OUTPUT_PREFIX}draft.docx"))


def wrong_award(tmp):
    copy_gt(tmp)
    set_award(tmp, "Solstice Contract Interiors\n")


def two_awards(tmp):
    copy_gt(tmp)
    set_award(tmp, f"{AWARD['vendor']}\nNorthpeak Office Supply\n")


def missing_award(tmp):
    copy_gt(tmp)
    os.remove(os.path.join(tmp, AWARD_FILE))


def empty_award(tmp):
    copy_gt(tmp)
    set_award(tmp, "\n")


def missing_sheet(tmp):
    copy_gt(tmp)
    os.remove(os.path.join(tmp, sheet_name("KWE")))


def swapped_decisions(tmp):
    copy_gt(tmp)
    for code, decision in (("ASF", "Compliant - Not Selected"), ("NPS", "Award")):
        paras, rows = gt_parts(code)
        rows[1][2] = decision
        write_sheet(tmp, code, paras, rows)
    set_award(tmp, "Northpeak Office Supply\n")


WINNER = AWARD["code"]

# ----------------------------------------------------------------- positives
case("exact ground truth", True, copy_gt)
case("no formatting, one paragraph per line, 'USD', '6 of 6', lower case, reasons reordered", True, messy_but_correct)
case("award.txt with the registered name and blank lines", True, registered_name_award)
case("award.txt with name and code", True, name_and_code_award)
case("cost without thousands separator, three decimals", True,
     modified(WINNER, rows=with_data_row(cost=VENDORS[WINNER]["total_usd"] + "0")))
case("extra unrelated files in the workspace", True, extra_files)
case("table with an extra empty row", True, modified(WINNER, rows=lambda r: r + [["", "", "", ""]]))
case("heading written as a normal paragraph", True, modified(WINNER, paragraphs=lambda p: p))

# ----------------------------------------------------------------- negatives
case("missing sheet (KWE)", False, missing_sheet)
case("ASF summary pasted from the superseded revision", False,
     modified("ASF", paragraphs=lambda p: p[:IDX_QUOTATION_BLOCK] + [quotation_block(2)] + p[IDX_QUOTATION_BLOCK + 1:]))
case("ASF cost replaced by BVL's total (44,250.00)", False, modified("ASF", rows=with_data_row(cost="44,250.00")))
case("ASF criteria count 5/6", False, modified("ASF", rows=with_data_row(met="5/6")))
case("award decision moved from ASF to NPS", False, swapped_decisions)
case("BVL with only one of its two reasons", False, modified("BVL", rows=with_data_row(reasons="Warranty below 5 years")))
case("GRD treated as compliant (due diligence ignored)", False,
     modified("GRD", rows=with_data_row(met="6/6", decision="Compliant - Not Selected", reasons="None")))
case("KWE total left in EUR (43,275.00)", False, modified("KWE", rows=with_data_row(cost="43,275.00")))
case("HMP volume discount wrongly applied (42,115.00)", False, modified("HMP", rows=with_data_row(cost="42,115.00")))
case("TRV total computed on 150 units (44,745.50)", False, modified("TRV", rows=with_data_row(cost="44,745.50")))
case("extra closing paragraph", False,
     modified(WINNER, paragraphs=lambda p: p + ["Prepared by the procurement assistant on 22 September 2025."]))
case("'Quotation 10' label pasted into the summary", False,
     modified(WINNER, paragraphs=lambda p: p[:IDX_QUOTATION_BLOCK] + ["Quotation 10"] + p[IDX_QUOTATION_BLOCK:]))
case("vendor name taken from the registered name", False,
     modified(WINNER, paragraphs=lambda p: p[:IDX_VENDOR_NAME] + [f"Vendor Name: {VENDORS[WINNER]['registered_name']}"] + p[IDX_VENDOR_NAME + 1:]))
case("no Evaluation Conclusion table", False, modified(WINNER, rows=lambda r: None))
case("award.txt names SOL (expired quote)", False, wrong_award)
case("award.txt lists two vendors", False, two_awards)
case("award.txt missing", False, missing_award)
case("award.txt empty", False, empty_award)

print()
print(f"{sum(results)}/{len(results)} cases behaved as expected")
sys.exit(0 if all(results) else 1)
