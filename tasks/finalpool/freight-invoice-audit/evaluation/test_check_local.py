#!/usr/bin/env python3
"""Self-test for evaluation/check_local.py.

Builds throw-away agent workspaces from the ground-truth workbook, mutates
them in the ways a real agent tends to get wrong, and asserts that the checker
accepts exactly the correct variants.

Run from the repository root or the task directory:
    python tasks/finalpool/freight-invoice-audit/evaluation/test_check_local.py
"""

import os
import shutil
import sys
import tempfile

import pandas as pd
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from check_local import OUTPUT_FILENAME, check_local  # noqa: E402

TASK_DIR = os.path.dirname(HERE)
GT_DIR = os.path.join(TASK_DIR, "groundtruth_workspace")
GT_FILE = os.path.join(GT_DIR, OUTPUT_FILENAME)

AUDIT, SUMMARY, DISPUTES = "Shipment Audit", "Service Summary", "Disputes"
COL = {name: i + 1 for i, name in enumerate([
    "Invoice Line", "Tracking Number", "Service", "Expected Zone", "Expected Billed Weight",
    "Expected Total", "Billed Total", "Variance", "Status", "Issues"])}

results = []


def case(name, expect_pass, builder):
    tmp = tempfile.mkdtemp(prefix="fia_")
    try:
        builder(tmp)
        ok, msg = check_local(tmp, GT_DIR)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    status = "ok " if ok == expect_pass else "BAD"
    results.append(ok == expect_pass)
    print(f"[{status}] {name:<52} -> {'pass' if ok else 'fail'}: {msg[:120]}")


def copy_gt(tmp):
    shutil.copy(GT_FILE, os.path.join(tmp, OUTPUT_FILENAME))
    return os.path.join(tmp, OUTPUT_FILENAME)


def mutate(fn):
    """Return a builder that copies the GT and applies fn(workbook)."""
    def build(tmp):
        path = copy_gt(tmp)
        wb = load_workbook(path)
        fn(wb)
        wb.save(path)
    return build


def first_row_with(ws, column, predicate):
    for r in range(2, ws.max_row + 1):
        if predicate(ws.cell(r, COL[column]).value):
            return r
    raise AssertionError(f"no row satisfies predicate on {column}")


# ----------------------------------------------------------------- positives
case("exact ground truth", True, copy_gt)


def via_pandas(tmp):
    sheets = pd.read_excel(GT_FILE, sheet_name=None)
    with pd.ExcelWriter(os.path.join(tmp, OUTPUT_FILENAME), engine="openpyxl") as w:
        for name, df in sheets.items():
            df.to_excel(w, sheet_name=name, index=False)


case("rewritten with pandas (no styling, blanks as NaN)", True, via_pandas)


def numbers_as_text(wb):
    ws = wb[AUDIT]
    for r in range(2, ws.max_row + 1):
        for c in (COL["Expected Total"], COL["Billed Total"], COL["Variance"]):
            if ws.cell(r, c).value is not None:
                ws.cell(r, c).value = f"{ws.cell(r, c).value}"


case("money cells stored as text", True, mutate(numbers_as_text))


def within_tolerance(wb):
    ws = wb[AUDIT]
    ws.cell(2, COL["Billed Total"]).value = float(ws.cell(2, COL["Billed Total"]).value) + 0.004


case("value off by 0.004 (inside tolerance)", True, mutate(within_tolerance))


def extra_sheet(wb):
    ws = wb.create_sheet("Scratch")
    ws.append(["notes"])


case("extra scratch sheet", True, mutate(extra_sheet))


def integer_as_float(wb):
    ws = wb[AUDIT]
    r = first_row_with(ws, "Expected Zone", lambda v: v is not None)
    ws.cell(r, COL["Expected Zone"]).value = float(ws.cell(r, COL["Expected Zone"]).value)


case("Expected Zone written as 4.0 instead of 4", True, mutate(integer_as_float))


def whitespace_in_text(wb):
    ws = wb[AUDIT]
    ws.cell(2, COL["Status"]).value = "  " + ws.cell(2, COL["Status"]).value + " "
    ws.cell(1, 1).value = "Invoice Line "


case("stray whitespace in text cells/header", True, mutate(whitespace_in_text))


def issues_without_spaces(wb):
    ws = wb[AUDIT]
    r = first_row_with(ws, "Issues", lambda v: v and ";" in v)
    ws.cell(r, COL["Issues"]).value = ws.cell(r, COL["Issues"]).value.replace("; ", ";")


case("Issues separated by ';' without a space", True, mutate(issues_without_spaces))

# ----------------------------------------------------------------- negatives
case("output file missing", False, lambda tmp: None)


def wrong_name(tmp):
    shutil.copy(GT_FILE, os.path.join(tmp, "freight_audit.xlsx"))


case("wrong file name", False, wrong_name)


def template_only(tmp):
    shutil.copy(os.path.join(TASK_DIR, "initial_workspace", "freight_audit_template.xlsx"),
                os.path.join(tmp, OUTPUT_FILENAME))


case("template saved without data", False, template_only)


def drop_sheet(wb):
    wb.remove(wb[DISPUTES])


case("Disputes sheet missing", False, mutate(drop_sheet))


def rename_header(wb):
    wb[SUMMARY].cell(1, 2).value = "Lines"


case("header renamed", False, mutate(rename_header))


def swap_columns(wb):
    ws = wb[AUDIT]
    a, b = COL["Expected Total"], COL["Billed Total"]
    for r in range(1, ws.max_row + 1):
        ws.cell(r, a).value, ws.cell(r, b).value = ws.cell(r, b).value, ws.cell(r, a).value


case("two columns swapped", False, mutate(swap_columns))


def off_by_cent(wb):
    ws = wb[AUDIT]
    r = first_row_with(ws, "Expected Total", lambda v: v is not None)
    ws.cell(r, COL["Expected Total"]).value = float(ws.cell(r, COL["Expected Total"]).value) + 0.01


case("Expected Total off by one cent", False, mutate(off_by_cent))


def zone_off_by_one(wb):
    ws = wb[AUDIT]
    r = first_row_with(ws, "Expected Zone", lambda v: v is not None)
    ws.cell(r, COL["Expected Zone"]).value = int(ws.cell(r, COL["Expected Zone"]).value) + 1


case("Expected Zone off by one", False, mutate(zone_off_by_one))


def status_spelling(wb):
    ws = wb[AUDIT]
    r = first_row_with(ws, "Status", lambda v: v == "Overbilled")
    ws.cell(r, COL["Status"]).value = "Over-billed"


case("Status spelled differently", False, mutate(status_spelling))


def issues_reordered(wb):
    ws = wb[AUDIT]
    r = first_row_with(ws, "Issues", lambda v: v and v.count(";") >= 2)
    parts = ws.cell(r, COL["Issues"]).value.split("; ")
    ws.cell(r, COL["Issues"]).value = "; ".join(reversed(parts))


case("Issues listed in a different order", False, mutate(issues_reordered))


def blank_filled_with_zero(wb):
    ws = wb[AUDIT]
    r = first_row_with(ws, "Status", lambda v: v == "Duplicate")
    ws.cell(r, COL["Expected Total"]).value = 0


case("empty expected cell of a duplicate filled with 0", False, mutate(blank_filled_with_zero))


def sort_by_tracking(wb):
    ws = wb[AUDIT]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    rows.sort(key=lambda r: r[1])
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j).value = v


case("Shipment Audit sorted by tracking number", False, mutate(sort_by_tracking))


def disputes_ascending(wb):
    ws = wb[DISPUTES]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    rows.reverse()
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j).value = v


case("Disputes sorted ascending", False, mutate(disputes_ascending))


def missing_row(wb):
    wb[AUDIT].delete_rows(10)


case("one audit row dropped", False, mutate(missing_row))


def extra_dispute(wb):
    ws = wb[DISPUTES]
    ws.append([9999, "MPS000000000000", "Ground", "Overbilled", "Fuel", 0.50])


case("extra dispute row", False, mutate(extra_dispute))


def wrong_count(wb):
    ws = wb[SUMMARY]
    ws.cell(2, 10).value = int(ws.cell(2, 10).value) + 1


case("Unbilled Shipments off by one", False, mutate(wrong_count))


def formulas_only(wb):
    ws = wb[AUDIT]
    r = first_row_with(ws, "Variance", lambda v: v is not None)
    ws.cell(r, COL["Variance"]).value = f"=G{r}-F{r}"


case("Variance written as a formula", False, mutate(formulas_only))

print()
print(f"{sum(results)}/{len(results)} cases behaved as expected")
sys.exit(0 if all(results) else 1)
