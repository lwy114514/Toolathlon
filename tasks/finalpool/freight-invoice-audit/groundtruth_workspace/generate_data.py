#!/usr/bin/env python3
"""Generate the agent-facing input files for the freight-invoice-audit task.

Fully deterministic: a single ``random.Random(SEED)`` instance drives every
choice, no set iteration order is relied upon, and no third-party RNG is used,
so re-running the script always reproduces byte-identical CSV content and
identical workbook cell values.

Outputs (written into ../initial_workspace/):
  shipment_manifest.csv          - ~620 shipments exported from our shipping system
  carrier_invoice_2025-09.csv    - the carrier's invoice lines for September 2025
  rate_card_2025.xlsx            - contract: base rates, zones, weekly fuel surcharge, terms
  freight_audit_template.xlsx    - the empty three-sheet audit workbook the agent fills

Noise / edge cases deliberately planted (all of them are covered by
initial_workspace/audit_rules.md):
  * invoice tracking numbers with stray whitespace / lower-case letters
  * destination postal codes with leading zeros and ZIP+4 suffixes
  * duplicate invoice lines (same tracking number billed twice)
  * invoice lines for tracking numbers that are not ours
  * September shipments the carrier never billed, plus late-August shipments
    carried over onto this invoice and October shipments that belong elsewhere
  * invoice ship dates later than the manifest ship date, some crossing into
    the next fuel-surcharge week
  * packages of exactly 1,728 cubic inches (no dimensional weight) next to
    slightly larger ones; whole-pound and sub-pound actual weights
  * declared values of exactly 100.00 (no surcharge) and above
  * carrier errors: wrong zone (both directions), billed weight one pound too
    high, dimensional weight ignored, service upgraded, residential surcharge
    on commercial addresses or missing on residential ones, fuel surcharge
    taken from the previous week, declared value surcharge charged at exactly
    100.00 or missing, and base charges off by 0.99 / 1.00 / 1.01 around the
    dispute threshold

Usage:  python generate_data.py
"""

import csv
import datetime as dt
import os
import random
import sys
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import audit_schema as S  # noqa: E402

SEED = 20250930

HERE = os.path.dirname(os.path.abspath(__file__))
INITIAL_WORKSPACE = os.path.join(os.path.dirname(HERE), "initial_workspace")

# --------------------------------------------------------------------------- #
# Plan
# --------------------------------------------------------------------------- #
N_SEPT_SHIPMENTS = 600
AUG_MANIFEST_DATES = ([dt.date(2025, 8, 27)] * 3 + [dt.date(2025, 8, 28)] * 3
                      + [dt.date(2025, 8, 29)] * 3 + [dt.date(2025, 8, 30)] * 3)
OCT_MANIFEST_DATES = [dt.date(2025, 10, 1)] * 4 + [dt.date(2025, 10, 2)] * 3 + [dt.date(2025, 10, 3)] * 3
N_AUG_BILLED = 5                 # late-August shipments carried over onto this invoice
N_UNBILLED_SEPT = 14             # September shipments the carrier forgot to bill
N_DUPLICATES = 9
N_UNKNOWN_LINES = 6              # somebody else's shipments on our invoice
N_TRACKING_NOISE = 24
WEEKDAY_DATE_SHIFT_FRACTION = 0.08

# carrier error type -> number of invoice lines (each line carries one error)
ERROR_PLAN = [
    ("zone_up", 12),
    ("zone_down", 4),
    ("weight_up", 8),
    ("dim_ignored", 8),
    ("service_up", 6),
    ("res_added", 10),
    ("res_missing", 4),
    ("fuel_prev_week", 10),
    ("declared_100", 4),
    ("declared_missing", 3),
    ("base_tweak_99", 1),
    ("base_tweak_100", 1),
    ("base_tweak_101", 1),
]

SERVICE_MIX = [("Ground", 68), ("Express Saver", 22), ("Priority Overnight", 10)]
ZONE_MIX = [(2, 8), (3, 17), (4, 22), (5, 20), (6, 15), (7, 8), (8, 10)]
DECLARED_CHOICES_CENTS = [3500, 4999, 7900, 10000, 12000, 15000, 19900, 25000, 30000, 37500, 48000, 60000]

# list price (cents) = a + b*w + (zone-2)*(c + d*w), rounded to the nearest 0.20.
# Every list price is therefore a multiple of 20 cents, so 20 % / 25 % / 30 %
# discounts are exact cent amounts and "round to the cent" never breaks a tie.
RATE_FORMULA = {
    "Ground": (860, 42, 120, 3),
    "Express Saver": (1420, 95, 240, 5),
    "Priority Overnight": (2680, 160, 380, 8),
}

FUEL_BP = {monday: int(Decimal(pct) * 100) for monday, pct in S.FUEL_WEEKS}


# --------------------------------------------------------------------------- #
# Small helpers
# --------------------------------------------------------------------------- #
def ceil_div(a, b):
    return -(-a // b)


def weighted_choice(rng, pairs):
    total = sum(w for _, w in pairs)
    r = rng.random() * total
    acc = 0
    for value, weight in pairs:
        acc += weight
        if r < acc:
            return value
    return pairs[-1][0]


def monday_of(date):
    return date - dt.timedelta(days=date.weekday())


def fuel_bp(date):
    return FUEL_BP[monday_of(date)]


def zone_of(postal):
    prefix = postal.strip()[:3]
    for lo, hi, zone in S.ZONE_RANGES:
        if lo <= prefix <= hi:
            return zone
    raise ValueError(f"postal prefix {prefix} is not covered by the zone table")


def money(cents):
    return f"{cents // 100}.{cents % 100:02d}"


# --------------------------------------------------------------------------- #
# Rate card and contract arithmetic (integer cents throughout)
# --------------------------------------------------------------------------- #
def build_rates():
    rates = {}
    for service, (a, b, c, d) in RATE_FORMULA.items():
        prev = None
        for w in range(1, S.MAX_WEIGHT_LB + 1):
            row = {}
            for z in S.ZONES:
                raw = a + b * w + (z - 2) * (c + d * w)
                row[z] = (raw + 10) // 20 * 20
            assert all(row[z] < row[z + 1] for z in S.ZONES[:-1]), (service, w)
            if prev is not None:
                assert all(prev[z] < row[z] for z in S.ZONES), (service, w)
            rates[(service, w)] = row
            prev = row
    return rates


RATES = build_rates()


def discounted_base(service, weight, zone):
    list_cents = RATES[(service, weight)][zone]
    factor = 100 - S.DISCOUNT_PERCENT[service]
    assert list_cents * factor % 100 == 0
    return list_cents * factor // 100


def rounded_actual(weight_tenths):
    return max(S.MIN_BILLABLE_WEIGHT_LB, ceil_div(weight_tenths, 10))


def dim_weight(length, width, height):
    volume = length * width * height
    return ceil_div(volume, S.DIM_DIVISOR) if volume > S.DIM_VOLUME_THRESHOLD else 0


def billable_weight(weight_tenths, length, width, height):
    return max(rounded_actual(weight_tenths), dim_weight(length, width, height))


def declared_surcharge(declared_cents):
    if declared_cents <= S.DECLARED_VALUE_THRESHOLD_CENTS:
        return 0
    return ceil_div(declared_cents, 10000) * S.DECLARED_VALUE_RATE_CENTS


def fuel_surcharge(base_cents, residential_cents, bp):
    return ((base_cents + residential_cents) * bp + 5000) // 10000


# --------------------------------------------------------------------------- #
# Shipments
# --------------------------------------------------------------------------- #
class Shipment:
    __slots__ = ("tracking", "ship_date", "service", "postal", "residential",
                 "weight_tenths", "length", "width", "height", "declared_cents", "order_id")

    def __init__(self, **kw):
        for k, v in kw.items():
            setattr(self, k, v)

    @property
    def zone(self):
        return zone_of(self.postal)

    def expected(self):
        weight = billable_weight(self.weight_tenths, self.length, self.width, self.height)
        base = discounted_base(self.service, weight, self.zone)
        res = S.RESIDENTIAL_SURCHARGE_CENTS if self.residential else 0
        decl = declared_surcharge(self.declared_cents)
        bp = fuel_bp(self.ship_date)
        fuel = fuel_surcharge(base, res, bp)
        return {"service": self.service, "zone": self.zone, "weight": weight, "base": base,
                "res": res, "decl": decl, "fuel": fuel, "bp": bp, "total": base + res + decl + fuel}

    def manifest_row(self):
        return [
            self.tracking,
            self.ship_date.isoformat(),
            self.service,
            self.postal,
            "Y" if self.residential else "N",
            f"{self.weight_tenths // 10}.{self.weight_tenths % 10}",
            self.length,
            self.width,
            self.height,
            money(self.declared_cents),
            self.order_id,
        ]


def sept_date_weights():
    pairs = []
    d = S.PERIOD_START + dt.timedelta(days=1)          # 2025-09-01 is Labor Day: no pickups
    while d <= S.PERIOD_END:
        if d.weekday() == 5:
            pairs.append((d, 1))
        elif d.weekday() < 5:
            pairs.append((d, 19))
        d += dt.timedelta(days=1)
    return pairs


def random_destination(rng):
    zone = weighted_choice(rng, ZONE_MIX)
    lo, hi, _ = rng.choice([r for r in S.ZONE_RANGES if r[2] == zone])
    prefix = rng.randint(int(lo), int(hi))
    postal = f"{prefix:03d}{rng.randint(0, 99):02d}"
    if rng.random() < 0.2:
        postal += f"-{rng.randint(1000, 9999)}"
    return postal


def random_package(rng):
    u = rng.random()
    if u < 0.45:
        tenths, dims = rng.randint(3, 50), (rng.randint(6, 14), rng.randint(4, 12), rng.randint(2, 8))
    elif u < 0.80:
        tenths, dims = rng.randint(51, 200), (rng.randint(10, 20), rng.randint(8, 16), rng.randint(4, 12))
    elif u < 0.95:
        tenths, dims = rng.randint(201, 450), (rng.randint(14, 26), rng.randint(12, 20), rng.randint(8, 16))
    else:
        tenths, dims = rng.randint(451, 700), (rng.randint(16, 28), rng.randint(14, 22), rng.randint(10, 18))
    length, width, height = sorted(dims, reverse=True)
    while billable_weight(tenths, length, width, height) > S.MAX_WEIGHT_LB:
        length -= 1
    return tenths, length, width, height


def random_declared(rng):
    return 0 if rng.random() < 0.5 else rng.choice(DECLARED_CHOICES_CENTS)


class IdFactory:
    def __init__(self, rng):
        self.rng = rng
        self.trackings = set()
        self.orders = set()

    def tracking(self):
        while True:
            t = f"MPS{self.rng.randint(10 ** 11, 10 ** 12 - 1)}"
            if t not in self.trackings:
                self.trackings.add(t)
                return t

    def order(self):
        while True:
            o = f"BW-{self.rng.randint(100000, 999999)}"
            if o not in self.orders:
                self.orders.add(o)
                return o


def make_shipment(rng, ids, ship_date):
    tenths, length, width, height = random_package(rng)
    return Shipment(
        tracking=ids.tracking(),
        ship_date=ship_date,
        service=weighted_choice(rng, SERVICE_MIX),
        postal=random_destination(rng),
        residential=rng.random() < 0.55,
        weight_tenths=tenths,
        length=length,
        width=width,
        height=height,
        declared_cents=random_declared(rng),
        order_id=ids.order(),
    )


def plant_special_packages(rng, shipments):
    """Overwrite a few random September shipments with boundary packages."""
    specials = [
        dict(weight_tenths=40, length=12, width=12, height=12),      # exactly 1,728 in^3: no dim weight
        dict(weight_tenths=40, length=12, width=12, height=12),
        dict(weight_tenths=40, length=13, width=12, height=12),      # 1,872 in^3 -> dim weight 14 lb
        dict(weight_tenths=40, length=13, width=12, height=12),
        dict(weight_tenths=50, length=10, width=8, height=6),        # whole-pound actual weights stay
        dict(weight_tenths=120, length=12, width=10, height=8),
        dict(weight_tenths=300, length=14, width=12, height=10),
        dict(weight_tenths=4, length=8, width=6, height=2),          # 0.4 lb -> minimum 1 lb
        dict(weight_tenths=4, length=9, width=6, height=3),
        dict(weight_tenths=696, length=16, width=12, height=8),      # 69.6 lb -> 70 lb (table maximum)
        dict(declared_cents=10000),                                  # exactly 100.00: no surcharge
        dict(declared_cents=10000),
        dict(declared_cents=10000),
        dict(declared_cents=10000),
        dict(declared_cents=10000),
        dict(declared_cents=10000),
        dict(declared_cents=15000),                                  # 2 x 1.05
        dict(declared_cents=25000),                                  # 3 x 1.05
        dict(declared_cents=60000),                                  # 6 x 1.05
    ]
    targets = rng.sample(range(len(shipments)), len(specials))
    for idx, override in zip(targets, specials):
        for k, v in override.items():
            setattr(shipments[idx], k, v)
        assert billable_weight(shipments[idx].weight_tenths, shipments[idx].length,
                               shipments[idx].width, shipments[idx].height) <= S.MAX_WEIGHT_LB


def build_manifest(rng, ids):
    date_pairs = sept_date_weights()
    sept = [make_shipment(rng, ids, weighted_choice(rng, date_pairs)) for _ in range(N_SEPT_SHIPMENTS)]
    plant_special_packages(rng, sept)
    aug = [make_shipment(rng, ids, d) for d in AUG_MANIFEST_DATES]
    oct_ = [make_shipment(rng, ids, d) for d in OCT_MANIFEST_DATES]
    return sept, aug, oct_


# --------------------------------------------------------------------------- #
# Invoice
# --------------------------------------------------------------------------- #
class InvoiceLine:
    __slots__ = ("tracking_text", "ship_date", "service", "zone", "weight", "base", "res", "decl",
                 "fuel", "shipment", "kind", "line_no")

    def __init__(self, tracking_text, ship_date, service, zone, weight, base, res, decl, fuel, shipment, kind):
        self.tracking_text = tracking_text
        self.ship_date = ship_date
        self.service = service
        self.zone = zone
        self.weight = weight
        self.base = base
        self.res = res
        self.decl = decl
        self.fuel = fuel
        self.shipment = shipment
        self.kind = kind
        self.line_no = None

    @classmethod
    def from_shipment(cls, shipment):
        e = shipment.expected()
        return cls(shipment.tracking, shipment.ship_date, e["service"], e["zone"], e["weight"],
                   e["base"], e["res"], e["decl"], e["fuel"], shipment, "clean")

    @property
    def total(self):
        return self.base + self.res + self.decl + self.fuel

    @property
    def variance(self):
        return self.total - self.shipment.expected()["total"] if self.shipment else None

    def copy(self, kind):
        c = InvoiceLine(self.tracking_text, self.ship_date, self.service, self.zone, self.weight,
                        self.base, self.res, self.decl, self.fuel, self.shipment, kind)
        return c

    def row(self):
        return [self.line_no, self.tracking_text, self.ship_date.isoformat(), self.service, self.zone,
                self.weight, money(self.base), money(self.res), money(self.decl), money(self.fuel),
                money(self.total)]


def eligible(line, error):
    s = line.shipment
    e = s.expected()
    if error == "zone_up":
        return e["zone"] < 8
    if error == "zone_down":
        return e["zone"] > 2
    if error == "weight_up":
        return e["weight"] < S.MAX_WEIGHT_LB
    if error == "dim_ignored":
        return e["weight"] > rounded_actual(s.weight_tenths)
    if error == "service_up":
        return s.service == "Ground"
    if error == "res_added":
        return not s.residential
    if error == "res_missing":
        return s.residential
    if error == "fuel_prev_week":
        return (monday_of(s.ship_date) - dt.timedelta(days=7)) in FUEL_BP
    if error == "declared_100":
        return s.declared_cents == S.DECLARED_VALUE_THRESHOLD_CENTS
    if error == "declared_missing":
        return s.declared_cents > S.DECLARED_VALUE_THRESHOLD_CENTS
    if error.startswith("base_tweak_"):
        return True
    raise ValueError(error)


def apply_error(line, error):
    s = line.shipment
    e = s.expected()
    bp = e["bp"]
    if error == "zone_up" or error == "zone_down":
        line.zone = e["zone"] + (1 if error == "zone_up" else -1)
        line.base = discounted_base(e["service"], e["weight"], line.zone)
        line.fuel = fuel_surcharge(line.base, line.res, bp)
    elif error == "weight_up":
        line.weight = e["weight"] + 1
        line.base = discounted_base(e["service"], line.weight, e["zone"])
        line.fuel = fuel_surcharge(line.base, line.res, bp)
    elif error == "dim_ignored":
        line.weight = rounded_actual(s.weight_tenths)
        line.base = discounted_base(e["service"], line.weight, e["zone"])
        line.fuel = fuel_surcharge(line.base, line.res, bp)
    elif error == "service_up":
        line.service = "Express Saver"
        line.base = discounted_base(line.service, e["weight"], e["zone"])
        line.fuel = fuel_surcharge(line.base, line.res, bp)
    elif error == "res_added":
        line.res = S.RESIDENTIAL_SURCHARGE_CENTS
        line.fuel = fuel_surcharge(line.base, line.res, bp)
    elif error == "res_missing":
        line.res = 0
        line.fuel = fuel_surcharge(line.base, line.res, bp)
    elif error == "fuel_prev_week":
        prev_bp = FUEL_BP[monday_of(s.ship_date) - dt.timedelta(days=7)]
        assert prev_bp != bp
        line.fuel = fuel_surcharge(line.base, line.res, prev_bp)
    elif error == "declared_100":
        line.decl = S.DECLARED_VALUE_RATE_CENTS
    elif error == "declared_missing":
        line.decl = 0
    elif error.startswith("base_tweak_"):
        line.base = e["base"] + int(error.rsplit("_", 1)[1])
    else:
        raise ValueError(error)
    line.kind = error
    assert line.variance != 0, (error, s.tracking)


def assign_errors(rng, lines):
    for error, count in ERROR_PLAN:
        candidates = [ln for ln in lines if ln.kind == "clean" and eligible(ln, error)]
        assert len(candidates) >= count, (error, len(candidates))
        for ln in rng.sample(candidates, count):
            apply_error(ln, error)


def apply_date_shifts(rng, lines):
    """The carrier prints its scan date, which can be after our ship date."""
    crossing = 0
    for ln in lines:
        d = ln.shipment.ship_date
        shift = 0
        if d.weekday() == 5:                       # Saturday pickup, scanned on Monday
            shift = 2
        elif d.weekday() == 4 and rng.random() < 0.25:
            shift = 3                              # Friday pickup, scanned on Monday
        elif rng.random() < WEEKDAY_DATE_SHIFT_FRACTION:
            shift = 1
        if shift:
            ln.ship_date = d + dt.timedelta(days=shift)
            if monday_of(ln.ship_date) != monday_of(d):
                crossing += 1
    return crossing


def make_unknown_line(rng, ids):
    date = weighted_choice(rng, sept_date_weights())
    service = weighted_choice(rng, SERVICE_MIX)
    zone = rng.randint(2, 8)
    weight = rng.randint(1, 40)
    base = discounted_base(service, weight, zone)
    res = S.RESIDENTIAL_SURCHARGE_CENTS if rng.random() < 0.5 else 0
    fuel = fuel_surcharge(base, res, fuel_bp(date))
    return InvoiceLine(ids.tracking(), date, service, zone, weight, base, res, 0, fuel, None, "unknown")


def mangle_tracking(rng, lines):
    variants = [
        lambda s: " " + s,
        lambda s: s + " ",
        lambda s: "  " + s + " ",
        lambda s: s.lower(),
        lambda s: s.lower() + " ",
        lambda s: " " + s.lower(),
    ]
    for ln in rng.sample(lines, N_TRACKING_NOISE):
        ln.tracking_text = rng.choice(variants)(ln.tracking_text)


def build_invoice(rng, ids, sept, aug):
    unbilled_idx = set(rng.sample(range(len(sept)), N_UNBILLED_SEPT))
    unbilled = [s for i, s in enumerate(sept) if i in unbilled_idx]
    billed = [s for i, s in enumerate(sept) if i not in unbilled_idx] + rng.sample(aug, N_AUG_BILLED)
    billed.sort(key=lambda s: (s.ship_date, s.tracking))

    lines = [InvoiceLine.from_shipment(s) for s in billed]
    assign_errors(rng, lines)
    crossing = apply_date_shifts(rng, lines)

    clean = [ln for ln in lines if ln.kind == "clean"]
    erroneous = [ln for ln in lines if ln.kind != "clean"]
    dup_sources = rng.sample(clean, N_DUPLICATES - 3) + rng.sample(erroneous, 3)
    for src in dup_sources:
        dup = src.copy("duplicate")
        pos = lines.index(src)
        lines.insert(rng.randint(pos + 1, len(lines)), dup)

    for _ in range(N_UNKNOWN_LINES):
        lines.insert(rng.randint(0, len(lines)), make_unknown_line(rng, ids))

    mangle_tracking(rng, lines)
    for i, ln in enumerate(lines, start=1):
        ln.line_no = i
    return lines, unbilled, crossing


# --------------------------------------------------------------------------- #
# Sanity checks
# --------------------------------------------------------------------------- #
def sanity_check(manifest, lines, unbilled):
    trackings = [s.tracking for s in manifest]
    assert len(set(trackings)) == len(trackings)
    for s in manifest:
        assert zone_of(s.postal) in S.ZONES
        assert 1 <= billable_weight(s.weight_tenths, s.length, s.width, s.height) <= S.MAX_WEIGHT_LB
        assert s.length >= s.width >= s.height >= 1

    by_tracking = {s.tracking: s for s in manifest}
    first_line = {}
    for ln in lines:
        canonical = ln.tracking_text.strip().upper()
        assert ln.line_no is not None
        if ln.kind == "unknown":
            assert canonical not in by_tracking
        else:
            assert canonical in by_tracking
        if ln.kind == "duplicate":
            assert canonical in first_line and first_line[canonical] < ln.line_no
        elif canonical not in first_line:
            first_line[canonical] = ln.line_no
        assert ln.total == ln.base + ln.res + ln.decl + ln.fuel

    billed_trackings = {ln.tracking_text.strip().upper() for ln in lines}
    for s in unbilled:
        assert s.tracking not in billed_trackings
    variances = sorted(ln.variance for ln in lines if ln.kind.startswith("base_tweak"))
    assert variances == [99, 100, 101], variances


# --------------------------------------------------------------------------- #
# Writers
# --------------------------------------------------------------------------- #
HEADER_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")


def style_header(ws, ncols):
    for col in range(1, ncols + 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)
    return path


def write_rate_card():
    wb = Workbook()
    ws = wb.active
    ws.title = S.SHEET_RATES
    ws.append(S.RATES_COLUMNS)
    for service in S.SERVICES:
        for w in range(1, S.MAX_WEIGHT_LB + 1):
            ws.append([service, w] + [RATES[(service, w)][z] / 100 for z in S.ZONES])
    for r in range(2, ws.max_row + 1):
        for c in range(3, 3 + len(S.ZONES)):
            ws.cell(row=r, column=c).number_format = "0.00"
    style_header(ws, len(S.RATES_COLUMNS))
    ws.column_dimensions["A"].width = 20

    ws = wb.create_sheet(S.SHEET_ZONES)
    ws.append(S.ZONES_COLUMNS)
    for lo, hi, zone in sorted(S.ZONE_RANGES):
        ws.append([lo, hi, zone])
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1).number_format = "@"
        ws.cell(row=r, column=2).number_format = "@"
    style_header(ws, len(S.ZONES_COLUMNS))

    ws = wb.create_sheet(S.SHEET_FUEL)
    ws.append(S.FUEL_COLUMNS)
    for monday, pct in S.FUEL_WEEKS:
        ws.append([monday.isoformat(), float(pct)])
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=2).number_format = "0.00"
    style_header(ws, len(S.FUEL_COLUMNS))
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 18

    ws = wb.create_sheet(S.SHEET_TERMS)
    ws.append(S.TERMS_COLUMNS)
    terms = [(S.TERM_DISCOUNT[s], S.DISCOUNT_PERCENT[s]) for s in S.SERVICES] + [
        (S.TERM_RESIDENTIAL, S.RESIDENTIAL_SURCHARGE_CENTS / 100),
        (S.TERM_DECLARED_THRESHOLD, S.DECLARED_VALUE_THRESHOLD_CENTS / 100),
        (S.TERM_DECLARED_RATE, S.DECLARED_VALUE_RATE_CENTS / 100),
        (S.TERM_DIM_DIVISOR, S.DIM_DIVISOR),
        (S.TERM_DIM_VOLUME, S.DIM_VOLUME_THRESHOLD),
        (S.TERM_MIN_WEIGHT, S.MIN_BILLABLE_WEIGHT_LB),
        (S.TERM_DISPUTE_THRESHOLD, S.DISPUTE_THRESHOLD_CENTS / 100),
        ("Origin ZIP", "43219"),
        ("Invoice number", S.INVOICE_NUMBER),
    ]
    for term, value in terms:
        ws.append([term, value])
    style_header(ws, len(S.TERMS_COLUMNS))
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 20

    path = os.path.join(INITIAL_WORKSPACE, S.RATE_CARD_FILENAME)
    wb.save(path)
    return path


def write_template():
    wb = Workbook()
    wb.remove(wb.active)
    widths = {
        S.SHEET_AUDIT: [12, 18, 18, 13, 21, 14, 12, 10, 16, 34],
        S.SHEET_SUMMARY: [18, 13, 18, 12, 14, 17, 18, 14, 16, 18],
        S.SHEET_DISPUTES: [12, 18, 18, 16, 34, 16],
    }
    for sheet, columns in S.REPORT_SHEETS.items():
        ws = wb.create_sheet(sheet)
        ws.append(columns)
        style_header(ws, len(columns))
        for i, width in enumerate(widths[sheet]):
            ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].width = width
    path = os.path.join(INITIAL_WORKSPACE, S.TEMPLATE_FILENAME)
    wb.save(path)
    return path


# --------------------------------------------------------------------------- #
def main():
    rng = random.Random(SEED)
    ids = IdFactory(rng)
    os.makedirs(INITIAL_WORKSPACE, exist_ok=True)

    sept, aug, oct_ = build_manifest(rng, ids)
    lines, unbilled, crossing = build_invoice(rng, ids, sept, aug)
    manifest = sept + aug + oct_
    rng.shuffle(manifest)
    sanity_check(manifest, lines, unbilled)

    paths = [
        write_csv(os.path.join(INITIAL_WORKSPACE, S.MANIFEST_FILENAME), S.MANIFEST_COLUMNS,
                  [s.manifest_row() for s in manifest]),
        write_csv(os.path.join(INITIAL_WORKSPACE, S.INVOICE_FILENAME), S.INVOICE_COLUMNS,
                  [ln.row() for ln in lines]),
        write_rate_card(),
        write_template(),
    ]

    kinds = {}
    for ln in lines:
        kinds[ln.kind] = kinds.get(ln.kind, 0) + 1
    dim_rated = sum(1 for s in manifest if dim_weight(s.length, s.width, s.height) > rounded_actual(s.weight_tenths))
    print(f"seed                 : {SEED}")
    print(f"manifest rows        : {len(manifest)} ({len(sept)} September, {len(aug)} August, {len(oct_)} October)")
    print(f"invoice lines        : {len(lines)}")
    print("line kinds           : " + ", ".join(f"{k}={v}" for k, v in sorted(kinds.items())))
    print(f"unbilled September   : {len(unbilled)}   carried-over August lines: {N_AUG_BILLED}")
    print(f"tracking noise       : {N_TRACKING_NOISE}   invoice dates crossing a fuel week: {crossing}")
    print(f"dim-rated shipments  : {dim_rated}")
    for p in paths:
        print(f"wrote {os.path.relpath(p, os.path.dirname(HERE))}")


if __name__ == "__main__":
    main()
