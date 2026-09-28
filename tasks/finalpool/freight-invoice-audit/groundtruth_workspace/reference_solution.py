#!/usr/bin/env python3
"""Reference solution for the freight-invoice-audit task.

Implements ``initial_workspace/audit_rules.md`` literally and writes the
ground-truth workbook ``freight_audit_2025-09.xlsx`` next to this script.  The
evaluator compares the agent's workbook against that file.

Money is handled as integer cents (parsed through ``Decimal``), so the ground
truth carries no floating-point drift.  ``generate_data.py`` guarantees that
every discounted list price is an exact cent amount and that no fuel-surcharge
rounding ever lands on a tie, so "round to the cent" is unambiguous.

Usage:
    python reference_solution.py
    python reference_solution.py --initial_workspace DIR --output FILE
"""

import argparse
import csv
import datetime as dt
import os
import sys
from collections import defaultdict
from decimal import Decimal

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import audit_schema as S  # noqa: E402

DEFAULT_INITIAL_WORKSPACE = os.path.join(os.path.dirname(HERE), "initial_workspace")
DEFAULT_OUTPUT = os.path.join(HERE, S.OUTPUT_FILENAME)


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def cents(value):
    """Exact cents from a CSV string or a workbook number."""
    return int((Decimal(str(value).strip()) * 100).to_integral_value())


def ceil_div(a, b):
    return -(-a // b)


def round_half_up_div(numerator, denominator):
    """numerator / denominator rounded half up to an integer (both positive)."""
    return (2 * numerator + denominator) // (2 * denominator)


def normalize_tracking(text):
    return (text or "").strip().upper()


def parse_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return dt.date.fromisoformat(str(value).strip()[:10])


def monday_of(date):
    return date - dt.timedelta(days=date.weekday())


# --------------------------------------------------------------------------- #
# Loading
# --------------------------------------------------------------------------- #
def load_manifest(path):
    manifest = {}
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        if [h.strip() for h in reader.fieldnames] != S.MANIFEST_COLUMNS:
            raise ValueError(f"unexpected manifest header: {reader.fieldnames}")
        for row in reader:
            tracking = normalize_tracking(row["tracking_number"])
            if tracking in manifest:
                raise ValueError(f"duplicate manifest tracking {tracking}")
            manifest[tracking] = {
                "ship_date": parse_date(row["ship_date"]),
                "service": row["service"].strip(),
                "postal": row["dest_postal_code"].strip(),
                "residential": row["residential"].strip().upper() == "Y",
                "weight_tenths": int((Decimal(row["actual_weight_lb"]) * 10).to_integral_value()),
                "length": int(row["length_in"]),
                "width": int(row["width_in"]),
                "height": int(row["height_in"]),
                "declared_cents": cents(row["declared_value_usd"]),
            }
    return manifest


def load_invoice(path):
    lines = []
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        if [h.strip() for h in reader.fieldnames] != S.INVOICE_COLUMNS:
            raise ValueError(f"unexpected invoice header: {reader.fieldnames}")
        for row in reader:
            lines.append({
                "line": int(row["invoice_line"]),
                "tracking": normalize_tracking(row["tracking_number"]),
                "ship_date": parse_date(row["ship_date"]),
                "service": row["service"].strip(),
                "zone": int(row["zone"]),
                "weight": int(row["billed_weight_lb"]),
                "base": cents(row["base_charge"]),
                "res": cents(row["residential_surcharge"]),
                "decl": cents(row["declared_value_surcharge"]),
                "fuel": cents(row["fuel_surcharge"]),
                "total": cents(row["total_charge"]),
            })
    lines.sort(key=lambda ln: ln["line"])
    return lines


def load_rate_card(path):
    wb = load_workbook(path, read_only=True, data_only=True)

    ws = wb[S.SHEET_RATES]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]
    if header != S.RATES_COLUMNS:
        raise ValueError(f"unexpected Base Rates header: {header}")
    rates = {}
    for row in rows:
        if row[0] is None:
            continue
        rates[(str(row[0]).strip(), int(row[1]))] = {z: cents(row[2 + i]) for i, z in enumerate(S.ZONES)}

    ws = wb[S.SHEET_ZONES]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]
    if header != S.ZONES_COLUMNS:
        raise ValueError(f"unexpected Zones header: {header}")
    zones = [(str(r[0]).strip().zfill(3), str(r[1]).strip().zfill(3), int(r[2])) for r in rows if r[0] is not None]

    ws = wb[S.SHEET_FUEL]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]
    if header != S.FUEL_COLUMNS:
        raise ValueError(f"unexpected Fuel Surcharge header: {header}")
    fuel_bp = {}
    for r in rows:
        if r[0] is None:
            continue
        week = parse_date(r[0])
        if week.weekday() != 0:
            raise ValueError(f"fuel week {week} does not start on a Monday")
        fuel_bp[week] = int((Decimal(str(r[1])) * 100).to_integral_value())   # percent -> basis points

    ws = wb[S.SHEET_TERMS]
    rows = ws.iter_rows(values_only=True)
    next(rows)
    terms = {str(r[0]).strip(): r[1] for r in rows if r[0] is not None}
    wb.close()

    contract = {
        "discount": {s: int(terms[S.TERM_DISCOUNT[s]]) for s in S.SERVICES},
        "residential": cents(terms[S.TERM_RESIDENTIAL]),
        "declared_threshold": cents(terms[S.TERM_DECLARED_THRESHOLD]),
        "declared_rate": cents(terms[S.TERM_DECLARED_RATE]),
        "dim_divisor": int(terms[S.TERM_DIM_DIVISOR]),
        "dim_volume": int(terms[S.TERM_DIM_VOLUME]),
        "min_weight": int(terms[S.TERM_MIN_WEIGHT]),
        "dispute_threshold": cents(terms[S.TERM_DISPUTE_THRESHOLD]),
    }
    return rates, zones, fuel_bp, contract


# --------------------------------------------------------------------------- #
# Rules
# --------------------------------------------------------------------------- #
def zone_of(postal, zones):
    prefix = postal[:3]
    for lo, hi, zone in zones:
        if lo <= prefix <= hi:
            return zone
    raise ValueError(f"postal prefix {prefix!r} not in zone table")


def billable_weight(m, contract):
    actual = max(contract["min_weight"], ceil_div(m["weight_tenths"], 10))            # §3.3 round up
    volume = m["length"] * m["width"] * m["height"]
    dim = ceil_div(volume, contract["dim_divisor"]) if volume > contract["dim_volume"] else 0
    return max(actual, dim)


def expected_charge(m, rates, zones, fuel_bp, contract):
    """§3: expected components of an audited line, in cents."""
    zone = zone_of(m["postal"], zones)
    weight = billable_weight(m, contract)
    list_price = rates[(m["service"], weight)][zone]
    base = round_half_up_div(list_price * (100 - contract["discount"][m["service"]]), 100)   # §3.4
    res = contract["residential"] if m["residential"] else 0                                # §3.5
    decl = 0                                                                                 # §3.6
    if m["declared_cents"] > contract["declared_threshold"]:
        decl = ceil_div(m["declared_cents"], 10000) * contract["declared_rate"]
    bp = fuel_bp[monday_of(m["ship_date"])]                                                  # §3.7 manifest date
    fuel = round_half_up_div((base + res) * bp, 10000)
    return {"service": m["service"], "zone": zone, "weight": weight, "base": base,
            "res": res, "decl": decl, "fuel": fuel, "total": base + res + decl + fuel}


def compute(manifest, invoice, rates, zones, fuel_bp, contract):
    audit_rows = []
    seen = set()
    for ln in invoice:
        tracking = ln["tracking"]
        row = {
            "Invoice Line": ln["line"],
            "Tracking Number": tracking,
            "Billed Total": ln["total"],
        }
        if tracking in seen:                                                    # §2.1
            m = manifest.get(tracking)
            row.update({"Service": m["service"] if m else ln["service"], "Expected Zone": None,
                        "Expected Billed Weight": None, "Expected Total": None,
                        "Variance": ln["total"], "Status": S.STATUS_DUPLICATE, "Issues": S.STATUS_DUPLICATE})
        elif tracking not in manifest:                                          # §2.2
            row.update({"Service": ln["service"], "Expected Zone": None, "Expected Billed Weight": None,
                        "Expected Total": None, "Variance": ln["total"],
                        "Status": S.STATUS_UNKNOWN, "Issues": S.STATUS_UNKNOWN})
        else:                                                                   # §2.3 / §3 / §4
            e = expected_charge(manifest[tracking], rates, zones, fuel_bp, contract)
            variance = ln["total"] - e["total"]
            issues = [code for code, differs in (
                (S.ISSUE_SERVICE, ln["service"] != e["service"]),
                (S.ISSUE_ZONE, ln["zone"] != e["zone"]),
                (S.ISSUE_WEIGHT, ln["weight"] != e["weight"]),
                (S.ISSUE_BASE, ln["base"] != e["base"]),
                (S.ISSUE_RESIDENTIAL, ln["res"] != e["res"]),
                (S.ISSUE_DECLARED, ln["decl"] != e["decl"]),
                (S.ISSUE_FUEL, ln["fuel"] != e["fuel"]),
            ) if differs]
            status = S.STATUS_OK if variance == 0 else (S.STATUS_OVER if variance > 0 else S.STATUS_UNDER)
            row.update({"Service": e["service"], "Expected Zone": e["zone"],
                        "Expected Billed Weight": e["weight"], "Expected Total": e["total"],
                        "Variance": variance, "Status": status,
                        "Issues": S.ISSUE_SEPARATOR.join(issues) if issues else S.NO_ISSUES})
        seen.add(tracking)
        audit_rows.append(row)

    dispute_rows = []                                                            # §5
    for row in audit_rows:
        if row["Status"] in (S.STATUS_DUPLICATE, S.STATUS_UNKNOWN):
            amount = row["Billed Total"]
        elif row["Status"] == S.STATUS_OVER and row["Variance"] >= contract["dispute_threshold"]:
            amount = row["Variance"]
        else:
            continue
        dispute_rows.append({c: row[c] for c in S.DISPUTE_COLUMNS if c != "Disputed Amount"} | {"Disputed Amount": amount})
    dispute_rows.sort(key=lambda r: (-r["Disputed Amount"], r["Invoice Line"]))

    billed_trackings = {ln["tracking"] for ln in invoice}
    unbilled = defaultdict(int)
    for tracking, m in manifest.items():
        if S.PERIOD_START <= m["ship_date"] <= S.PERIOD_END and tracking not in billed_trackings:
            unbilled[m["service"]] += 1

    summary_rows = []
    for service in sorted({r["Service"] for r in audit_rows}):
        rows = [r for r in audit_rows if r["Service"] == service]
        audited = [r for r in rows if r["Status"] in S.AUDITED_STATUSES]
        disputes = [d for d in dispute_rows if d["Service"] == service]
        summary_rows.append({
            "Service": service,
            "Invoice Lines": len(rows),
            "Audited Shipments": len(audited),
            "Billed Total": sum(r["Billed Total"] for r in rows),
            "Expected Total": sum(r["Expected Total"] for r in audited),
            "Overbilled Amount": sum(r["Variance"] for r in audited if r["Status"] == S.STATUS_OVER),
            "Underbilled Amount": sum(-r["Variance"] for r in audited if r["Status"] == S.STATUS_UNDER),
            "Disputed Lines": len(disputes),
            "Disputed Amount": sum(d["Disputed Amount"] for d in disputes),
            "Unbilled Shipments": unbilled.get(service, 0),
        })

    return {S.SHEET_AUDIT: audit_rows, S.SHEET_SUMMARY: summary_rows, S.SHEET_DISPUTES: dispute_rows}


# --------------------------------------------------------------------------- #
# Output
# --------------------------------------------------------------------------- #
MONEY_COLUMNS = {
    "Expected Total", "Billed Total", "Variance", "Overbilled Amount", "Underbilled Amount",
    "Disputed Amount",
}


def to_cell(column, value):
    if value is None:
        return None
    if column in MONEY_COLUMNS:
        return round(value / 100, 2)
    return value


def write_report(report, template_path, output_path):
    wb = load_workbook(template_path)
    for sheet, columns in S.REPORT_SHEETS.items():
        ws = wb[sheet]
        header = [str(c.value).strip() for c in ws[1][: len(columns)]]
        if header != columns:
            raise ValueError(f"template sheet {sheet!r} header mismatch: {header}")
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for row in report[sheet]:
            ws.append([to_cell(c, row[c]) for c in columns])
        for r in range(2, ws.max_row + 1):
            for c_idx, col in enumerate(columns, start=1):
                if col in MONEY_COLUMNS and report[sheet][r - 2][col] is not None:
                    ws.cell(row=r, column=c_idx).number_format = "0.00"
    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="Build the ground-truth freight audit workbook")
    parser.add_argument("--initial_workspace", default=DEFAULT_INITIAL_WORKSPACE)
    parser.add_argument("--output", default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    ws_dir = args.initial_workspace
    manifest = load_manifest(os.path.join(ws_dir, S.MANIFEST_FILENAME))
    invoice = load_invoice(os.path.join(ws_dir, S.INVOICE_FILENAME))
    rates, zones, fuel_bp, contract = load_rate_card(os.path.join(ws_dir, S.RATE_CARD_FILENAME))

    report = compute(manifest, invoice, rates, zones, fuel_bp, contract)
    write_report(report, os.path.join(ws_dir, S.TEMPLATE_FILENAME), args.output)

    audit = report[S.SHEET_AUDIT]
    by_status = defaultdict(int)
    for r in audit:
        by_status[r["Status"]] += 1
    print(f"manifest rows: {len(manifest)}, invoice lines: {len(invoice)}")
    print("statuses: " + ", ".join(f"{k}={v}" for k, v in sorted(by_status.items())))
    print(f"disputes: {len(report[S.SHEET_DISPUTES])}, disputed amount: "
          f"{sum(d['Disputed Amount'] for d in report[S.SHEET_DISPUTES]) / 100:,.2f}")
    for s in report[S.SHEET_SUMMARY]:
        print(f"  {s['Service']:<20} lines={s['Invoice Lines']:>3} audited={s['Audited Shipments']:>3} "
              f"billed={s['Billed Total'] / 100:>10,.2f} expected={s['Expected Total'] / 100:>10,.2f} "
              f"over={s['Overbilled Amount'] / 100:>8,.2f} under={s['Underbilled Amount'] / 100:>7,.2f} "
              f"disputes={s['Disputed Lines']:>2} unbilled={s['Unbilled Shipments']}")
    print(f"wrote {args.output}")


if __name__ == "__main__":
    main()
