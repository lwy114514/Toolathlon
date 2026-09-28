#!/usr/bin/env python3
"""Generate the agent-facing input files for the overtime-payroll task.

Fully deterministic: a single ``random.Random(SEED)`` instance drives every
choice, no set iteration order is relied upon, and no third-party RNG is used,
so re-running the script always reproduces byte-identical CSV content and
identical workbook cell values.

Outputs (written into ../initial_workspace/):
  employee_roster.xlsx           - 50 employees, 45 of them overtime-eligible
  attendance_export.csv          - ~1,100 raw badge punches with realistic noise
  holiday_calendar_2025.csv      - company holiday calendar for 2025
  overtime_report_template.xlsx  - the empty three-sheet report the agent fills

Noise / edge cases deliberately planted in the export (all of them are
covered by initial_workspace/overtime_policy.md):
  * exact duplicate rows (terminal re-sync), including a duplicated invalid row
  * employee IDs with stray whitespace / lower-case letters
  * rows outside the pay period (late July, early September)
  * night shifts crossing midnight on both period boundaries
  * invalid punches: missing clock-out, clock-out <= clock-in, > 16 h spans
  * an exactly-16-hour shift (valid) next to a 16 h 01 min one (invalid)
  * part-time shifts of exactly 6 h (no meal break) and 6 h 01-20 min (break)
  * split shifts whose two halves only exceed 8 h when added together
  * a company holiday on a Friday and another on a Saturday
  * punches for contractors, interns and terminated (non-roster) badges
  * eligible employees with no punches, or only invalid punches
  * employees engineered to land at 29.75 / 30.00 / 30.25 overtime hours

Usage:  python generate_data.py
"""

import csv
import datetime as dt
import os
import random
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import payroll_schema as S  # noqa: E402

SEED = 20250831

HERE = os.path.dirname(os.path.abspath(__file__))
INITIAL_WORKSPACE = os.path.join(os.path.dirname(HERE), "initial_workspace")

# --------------------------------------------------------------------------- #
# Static vocabularies
# --------------------------------------------------------------------------- #
FIRST_NAMES = [
    "Aaliyah", "Adrian", "Aiko", "Alejandro", "Amara", "Anders", "Anika", "Bartosz",
    "Beatriz", "Bianca", "Callum", "Carmen", "Chidi", "Dario", "Dmitri", "Elena",
    "Emeka", "Esme", "Farah", "Felix", "Fiona", "Gabriel", "Grace", "Hana", "Hugo",
    "Ines", "Isaac", "Ivana", "Jamal", "Jonas", "Julia", "Kai", "Kenji", "Layla",
    "Leon", "Liam", "Lucia", "Malik", "Marta", "Mateo", "Maya", "Mikael", "Nadia",
    "Nikolai", "Noor", "Olivia", "Omar", "Oscar", "Petra", "Priya", "Rafael",
    "Rania", "Rohan", "Rosa", "Samir", "Sana", "Sebastian", "Selin", "Sofia",
    "Tariq", "Teodor", "Thea", "Tomasz", "Uma", "Viktor", "Wei", "Yara", "Yusuf",
    "Zara", "Zoe",
]
LAST_NAMES = [
    "Abbott", "Adeyemi", "Almeida", "Andersen", "Baptiste", "Becker", "Bergstrom",
    "Brennan", "Castillo", "Chen", "Costa", "Dahl", "Delgado", "Dubois", "Eriksen",
    "Farouk", "Fischer", "Fontaine", "Garcia", "Haddad", "Hansen", "Hoffmann",
    "Ibrahim", "Iyer", "Jansen", "Kaur", "Kimura", "Kovac", "Kowalski", "Larsen",
    "Lindgren", "Lindqvist", "Lopez", "Marchetti", "Marin", "Mendes", "Moreau",
    "Nakamura", "Nguyen", "Novak", "Okafor", "Okoye", "Olsen", "Ortega", "Osei",
    "Park", "Patel", "Petrov", "Quinn", "Rahman", "Ramirez", "Rasmussen", "Reyes",
    "Rossi", "Santos", "Schneider", "Silva", "Sorensen", "Suzuki", "Takahashi",
    "Tanaka", "Torres", "Vargas", "Vasquez", "Vidal", "Volkov", "Weber", "Yamamoto",
    "Zhang", "Zielinski",
]

# Department -> (full-time count, part-time count)
DEPARTMENT_PLAN = [
    ("Machining", 9, 1),
    ("Assembly", 9, 2),
    ("Quality Assurance", 5, 1),
    ("Logistics", 5, 2),
    ("Maintenance", 5, 1),
    ("Administration", 4, 1),
]
CONTRACTOR_DEPARTMENTS = ["Machining", "Maintenance", "Administration"]
INTERN_DEPARTMENTS = ["Quality Assurance", "Administration"]

# Every rate below has a cent value divisible by 8.  Paid hours are always a
# multiple of 0.25 h and the multipliers are 1.0 / 1.5 / 2.0 / 2.5, so every
# pay figure is a multiple of rate/8 - i.e. exactly representable with two
# decimals.  "Round to 2 decimals" therefore never has to break a tie, which
# keeps the ground truth independent of the rounding mode an agent uses.
RATE_TABLE = {
    "Machining": [24.00, 26.40, 28.00, 28.80, 30.40, 31.20, 33.60],
    "Assembly": [21.60, 22.40, 23.20, 24.00, 25.60, 26.40],
    "Quality Assurance": [27.20, 28.80, 30.40, 32.00, 33.60],
    "Logistics": [20.80, 21.60, 22.40, 24.00, 25.60],
    "Maintenance": [28.00, 29.60, 31.20, 32.80, 35.20],
    "Administration": [22.40, 24.00, 26.40, 28.00, 30.40, 36.00],
}
PART_TIME_RATES = [17.60, 18.40, 19.20, 20.80, 21.60]
CONTRACTOR_RATES = [40.00, 44.80, 48.00, 52.00]
INTERN_RATES = [16.00, 17.60]

TERMINALS = ["GATE-A", "GATE-B", "SHOP-1", "SHOP-2", "ADMIN-1"]
HOME_TERMINAL = {
    "Machining": "SHOP-1",
    "Assembly": "SHOP-2",
    "Quality Assurance": "GATE-A",
    "Logistics": "GATE-B",
    "Maintenance": "SHOP-1",
    "Administration": "ADMIN-1",
}

HOLIDAYS_2025 = [
    ("2025-01-01", "New Year's Day"),
    ("2025-04-18", "Good Friday"),
    ("2025-05-26", "Memorial Day"),
    ("2025-07-04", "Independence Day"),
    ("2025-08-15", "Founders' Day"),
    ("2025-08-23", "Plant Anniversary"),
    ("2025-09-01", "Labor Day"),
    ("2025-11-27", "Thanksgiving Day"),
    ("2025-11-28", "Day after Thanksgiving"),
    ("2025-12-25", "Christmas Day"),
    ("2025-12-26", "Boxing Day"),
]
HOLIDAY_DATES = {dt.date.fromisoformat(d) for d, _ in HOLIDAYS_2025}

# Dates covered by the raw export: a little wider than the pay period.
EXPORT_START = dt.date(2025, 7, 30)
EXPORT_END = dt.date(2025, 9, 2)

# Terminated employees whose badges still produce punches (not in roster).
UNKNOWN_IDS = ["EMP-0987", "EMP-0993"]

WEEKDAY_P = {"day": 0.92, "night": 0.90, "part": 0.60, "split": 0.90, "heavy": 0.95}
WEEKEND_P = {"day": 0.08, "night": 0.15, "part": 0.10, "split": 0.12, "heavy": 0.35}
HOLIDAY_P = {"day": 0.30, "night": 0.45, "part": 0.25, "split": 0.35, "heavy": 0.60}

ALERT_TARGETS = [30.00, 30.25, 29.75]   # engineered boundary employees


# --------------------------------------------------------------------------- #
# Small helpers
# --------------------------------------------------------------------------- #
def minutes(h, m=0):
    return h * 60 + m


def at(date, minute_of_day):
    return dt.datetime.combine(date, dt.time(0, 0)) + dt.timedelta(minutes=minute_of_day)


def daterange(start, end):
    d = start
    while d <= end:
        yield d
        d += dt.timedelta(days=1)


def in_period(date):
    return S.PERIOD_START <= date <= S.PERIOD_END


def day_type(date):
    if date in HOLIDAY_DATES:
        return "holiday"
    if date.weekday() >= 5:
        return "weekend"
    return "weekday"


def paid_hours(raw_minutes):
    """Policy 5: unpaid 30-min break if raw > 6 h, then round to 0.25 h."""
    if raw_minutes > S.MEAL_BREAK_THRESHOLD_MINUTES:
        raw_minutes -= S.MEAL_BREAK_MINUTES
    return round(raw_minutes / S.ROUNDING_MINUTES) * (S.ROUNDING_MINUTES / 60.0)


def total_overtime(shifts):
    """Policy 6 applied to a list of valid (clock_in, clock_out) pairs."""
    per_day = defaultdict(float)
    weekend = holiday = 0.0
    for ci, co in shifts:
        h = paid_hours(int((co - ci).total_seconds() // 60))
        t = day_type(ci.date())
        if t == "weekday":
            per_day[ci.date()] += h
        elif t == "weekend":
            weekend += h
        else:
            holiday += h
    weekday_ot = sum(max(v - S.REGULAR_HOURS_PER_SHIFT, 0.0) for v in per_day.values())
    return weekday_ot + weekend + holiday


class Punch:
    __slots__ = ("employee_id", "clock_in", "clock_out", "terminal")

    def __init__(self, employee_id, clock_in, clock_out, terminal):
        self.employee_id = employee_id
        self.clock_in = clock_in
        self.clock_out = clock_out          # datetime or None
        self.terminal = terminal

    def row(self):
        return [
            self.employee_id,
            self.clock_in.strftime("%Y-%m-%d %H:%M"),
            self.clock_out.strftime("%Y-%m-%d %H:%M") if self.clock_out else "",
            self.terminal,
        ]

    def copy(self):
        return Punch(self.employee_id, self.clock_in, self.clock_out, self.terminal)


# --------------------------------------------------------------------------- #
# Roster
# --------------------------------------------------------------------------- #
def build_roster(rng):
    people = []
    for dept, n_ft, n_pt in DEPARTMENT_PLAN:
        people += [("Full-time", dept)] * n_ft
        people += [("Part-time", dept)] * n_pt
    people += [("Contractor", d) for d in CONTRACTOR_DEPARTMENTS]
    people += [("Intern", d) for d in INTERN_DEPARTMENTS]
    assert len(people) == 50

    rng.shuffle(people)  # interleave departments across the ID range

    name_pool = [(f, l) for f in FIRST_NAMES for l in LAST_NAMES]
    chosen_names = rng.sample(name_pool, len(people))

    roster = []
    for i, ((etype, dept), (first, last)) in enumerate(zip(people, chosen_names)):
        if etype == "Full-time":
            rate = rng.choice(RATE_TABLE[dept])
        elif etype == "Part-time":
            rate = rng.choice(PART_TIME_RATES)
        elif etype == "Contractor":
            rate = rng.choice(CONTRACTOR_RATES)
        else:
            rate = rng.choice(INTERN_RATES)
        assert round(rate * 100) % 8 == 0, rate
        hire = dt.date(2013, 1, 1) + dt.timedelta(days=rng.randint(0, 12 * 365 + 150))
        assert hire < S.PERIOD_START
        roster.append({
            "Employee ID": f"EMP-{1001 + i}",
            "Name": f"{first} {last}",
            "Department": dept,
            "Employment Type": etype,
            "Hourly Rate": rate,
            "Hire Date": hire.isoformat(),
        })
    return roster


def assign_profiles(roster):
    """Decide how each eligible employee behaves in the export."""
    profiles = {}
    counters = defaultdict(int)
    for emp in roster:
        if emp["Employment Type"] not in S.ELIGIBLE_EMPLOYMENT_TYPES:
            continue
        key = (emp["Department"], emp["Employment Type"])
        k = counters[key]
        counters[key] += 1
        eid = emp["Employee ID"]
        if emp["Employment Type"] == "Part-time":
            profiles[eid] = "absent" if key == ("Administration", "Part-time") and k == 0 else "part"
        elif emp["Department"] == "Maintenance" and k < 4:
            profiles[eid] = "night"
        elif emp["Department"] == "Logistics" and k < 2:
            profiles[eid] = "split"
        elif emp["Department"] == "Machining" and k < 2:
            profiles[eid] = "heavy"
        elif emp["Department"] == "Assembly" and k == 0:
            profiles[eid] = "heavy"
        elif emp["Department"] == "Assembly" and k == 1:
            profiles[eid] = "absent"          # on unpaid leave all month
        elif emp["Department"] == "Quality Assurance" and k == 0:
            profiles[eid] = "only_invalid"    # every punch of theirs is broken
        else:
            profiles[eid] = "day"
    return profiles


# --------------------------------------------------------------------------- #
# Shift generation
# --------------------------------------------------------------------------- #
def shift_specs(rng, profile):
    """Return a list of (start_minute_of_day, raw_duration_minutes)."""
    if profile == "day":
        start = minutes(6, 45) + rng.randint(0, 105)
        u = rng.random()
        if u < 0.70:
            raw = rng.randint(495, 540)      # 8h15 - 9h00 raw -> 0 to 0.5 h OT
        elif u < 0.85:
            raw = rng.randint(545, 600)      # 9h05 - 10h00 raw -> 0.5 to 1.5 h OT
        else:
            raw = rng.randint(430, 490)      # short day, no OT
        return [(start, raw)]
    if profile == "heavy":
        start = minutes(6, 0) + rng.randint(0, 60)
        return [(start, rng.randint(570, 660))]
    if profile == "night":
        start = minutes(21, 30) + rng.randint(0, 60)
        return [(start, rng.randint(510, 570))]
    if profile == "part":
        start = minutes(8, 0) + rng.randint(0, 300)
        u = rng.random()
        if u < 0.25:
            raw = 360                        # exactly 6 h -> no meal break
        elif u < 0.40:
            raw = rng.randint(361, 380)      # just over 6 h -> break applies
        else:
            raw = rng.randint(240, 355)
        return [(start, raw)]
    if profile == "split":
        s1 = minutes(5, 45) + rng.randint(0, 30)
        r1 = rng.randint(195, 255)
        s2 = minutes(13, 45) + rng.randint(0, 45)
        r2 = rng.randint(250, 345)
        return [(s1, r1), (s2, r2)]
    raise ValueError(profile)


def works_today(rng, profile, date):
    if date in HOLIDAY_DATES:
        p = HOLIDAY_P[profile]
    elif date.weekday() >= 5:
        p = WEEKEND_P[profile]
    else:
        p = WEEKDAY_P[profile]
    return rng.random() < p


def pick_terminal(rng, dept):
    home = HOME_TERMINAL.get(dept, "GATE-A")
    return home if rng.random() < 0.85 else rng.choice(TERMINALS)


def generate_bulk(rng, roster, profiles):
    punches = []
    used_dates = defaultdict(set)      # eid -> {date of clock_in}
    by_id = {e["Employee ID"]: e for e in roster}

    forced_night = {          # first night worker: boundary-crossing shifts
        dt.date(2025, 7, 31): (minutes(22, 0), 510),   # ends Aug 1 -> out of period
        dt.date(2025, 8, 14): (minutes(22, 0), 525),   # Thu -> Fri holiday: weekday
        dt.date(2025, 8, 15): (minutes(22, 10), 510),  # holiday night shift
        dt.date(2025, 8, 31): (minutes(22, 0), 510),   # Sun -> Labor Day: weekend
    }
    night_ids = [eid for eid in sorted(profiles) if profiles[eid] == "night"]
    first_night = night_ids[0]

    day_ids = [eid for eid in sorted(profiles) if profiles[eid] == "day"]
    part_ids = [eid for eid in sorted(profiles) if profiles[eid] == "part"]
    forced_saturday_holiday = set(day_ids[:3] + part_ids[:1])

    for eid in sorted(profiles):
        profile = profiles[eid]
        if profile in ("absent", "only_invalid"):
            continue
        dept = by_id[eid]["Department"]
        for date in daterange(EXPORT_START, EXPORT_END):
            specs = None
            if eid == first_night and date in forced_night:
                specs = [forced_night[date]]
            elif eid in forced_saturday_holiday and date == dt.date(2025, 8, 23):
                specs = shift_specs(rng, profile)
            elif works_today(rng, profile, date):
                specs = shift_specs(rng, profile)
            if not specs:
                continue
            for start, raw in specs:
                ci = at(date, start)
                punches.append(Punch(eid, ci, ci + dt.timedelta(minutes=raw), pick_terminal(rng, dept)))
            used_dates[eid].add(date)

    # Contractors and interns: normal looking day shifts, must be excluded.
    for emp in roster:
        if emp["Employment Type"] in S.ELIGIBLE_EMPLOYMENT_TYPES:
            continue
        eid = emp["Employee ID"]
        for date in daterange(EXPORT_START, EXPORT_END):
            if date.weekday() < 5 and date not in HOLIDAY_DATES and rng.random() < 0.8:
                start, raw = shift_specs(rng, "day")[0]
                ci = at(date, start)
                punches.append(Punch(eid, ci, ci + dt.timedelta(minutes=raw), pick_terminal(rng, emp["Department"])))
                used_dates[eid].add(date)

    # Terminated badges that are not in the roster at all.
    for eid in UNKNOWN_IDS:
        weekdays = [d for d in daterange(S.PERIOD_START, S.PERIOD_END) if d.weekday() < 5]
        for date in sorted(rng.sample(weekdays, 6)):
            start, raw = shift_specs(rng, "day")[0]
            ci = at(date, start)
            punches.append(Punch(eid, ci, ci + dt.timedelta(minutes=raw), rng.choice(TERMINALS)))
            used_dates[eid].add(date)

    return punches, used_dates


def free_dates(used, start, end, weekday_only=False, exclude_holidays=False, need_next_free=False):
    out = []
    for d in daterange(start, end):
        if d in used:
            continue
        if weekday_only and d.weekday() >= 5:
            continue
        if exclude_holidays and d in HOLIDAY_DATES:
            continue
        if need_next_free and (d + dt.timedelta(days=1)) in used:
            continue
        out.append(d)
    return out


def add_engineered_rows(rng, roster, profiles, punches, used_dates):
    """Plant the documented edge cases on top of the random bulk."""
    by_id = {e["Employee ID"]: e for e in roster}
    day_ids = [eid for eid in sorted(profiles) if profiles[eid] == "day"]
    contractor_ids = [e["Employee ID"] for e in roster if e["Employment Type"] == "Contractor"]
    only_invalid_id = [eid for eid in sorted(profiles) if profiles[eid] == "only_invalid"][0]

    pool = list(day_ids)      # each engineered case consumes a distinct day worker

    def claim(**kw):
        """Pop the first pooled employee that still has a free date matching kw."""
        for i, eid in enumerate(pool):
            cands = free_dates(used_dates[eid], S.PERIOD_START, S.PERIOD_END, **kw)
            if cands:
                pool.pop(i)
                d = cands[0]
                used_dates[eid].add(d)
                return eid, d
        raise RuntimeError(f"no pooled employee has a free date for {kw}")

    def add(eid, ci, co):
        punches.append(Punch(eid, ci, co, pick_terminal(rng, by_id[eid]["Department"])))

    # --- exactly 16 h (valid) and 16 h 01 min (invalid) --------------------
    eid, d = claim(weekday_only=True, exclude_holidays=True)
    add(eid, at(d, minutes(6, 0)), at(d, minutes(22, 0)))          # 960 min -> valid, 15.5 h
    eid, d = claim(weekday_only=True, exclude_holidays=True)
    add(eid, at(d, minutes(6, 0)), at(d, minutes(22, 1)))          # 961 min -> invalid

    # --- alert boundary employees: top up with weekend shifts ---------------
    valid_shifts = defaultdict(list)
    for p in punches:
        if p.employee_id in profiles and p.clock_out and in_period(p.clock_in.date()):
            raw = int((p.clock_out - p.clock_in).total_seconds() // 60)
            if 0 < raw <= 16 * 60:
                valid_shifts[p.employee_id].append((p.clock_in, p.clock_out))
    candidates = [eid for eid in pool if total_overtime(valid_shifts[eid]) <= 22.0]
    assert len(candidates) >= len(ALERT_TARGETS), candidates
    boundary = {}
    for eid, target in zip(candidates, ALERT_TARGETS):
        pool.remove(eid)
        needed = round(target - total_overtime(valid_shifts[eid]), 2)
        assert needed > 0
        weekend_free = free_dates(used_dates[eid], S.PERIOD_START, S.PERIOD_END, exclude_holidays=True)
        weekend_free = [d for d in weekend_free if d.weekday() >= 5]
        while needed > 0:
            chunk = min(needed, 12.0)
            raw = int(round(chunk * 60))
            if raw > S.MEAL_BREAK_THRESHOLD_MINUTES:
                raw += S.MEAL_BREAK_MINUTES
            date = weekend_free.pop(0)
            used_dates[eid].add(date)
            ci = at(date, minutes(8, 0) + rng.randint(0, 30))
            co = ci + dt.timedelta(minutes=raw)
            add(eid, ci, co)
            valid_shifts[eid].append((ci, co))
            needed = round(needed - chunk, 2)
        assert abs(total_overtime(valid_shifts[eid]) - target) < 1e-9
        boundary[eid] = target

    # --- missing clock-out ---------------------------------------------------
    missing_out_rows = []
    for _ in range(6):
        eid, d = claim(weekday_only=True)
        add(eid, at(d, minutes(7, 30) + rng.randint(0, 40)), None)
        missing_out_rows.append(punches[-1])

    # --- clock-out equal to clock-in ----------------------------------------
    for _ in range(2):
        eid, d = claim(weekday_only=True)
        t = at(d, minutes(7, 0) + rng.randint(0, 60))
        add(eid, t, t)

    # --- clock-out before clock-in (typed in reverse) -----------------------
    for _ in range(3):
        eid, d = claim()
        add(eid, at(d, minutes(14, 5)), at(d, minutes(9, 40)))

    # --- badge left open for more than 16 h ----------------------------------
    for _ in range(4):
        eid, d = claim(need_next_free=True)
        used_dates[eid].add(d + dt.timedelta(days=1))
        add(eid, at(d, minutes(7, 0)), at(d + dt.timedelta(days=1), minutes(9, 15)))

    # --- employee whose only punches are invalid ------------------------------
    d = free_dates(used_dates[only_invalid_id], S.PERIOD_START, S.PERIOD_END, weekday_only=True)[0]
    used_dates[only_invalid_id].add(d)
    add(only_invalid_id, at(d, minutes(8, 0)), None)
    d = free_dates(used_dates[only_invalid_id], S.PERIOD_START, S.PERIOD_END, weekday_only=True)[0]
    used_dates[only_invalid_id].add(d)
    add(only_invalid_id, at(d, minutes(15, 20)), at(d, minutes(8, 5)))

    # --- broken rows that must be ignored for *other* reasons ------------------
    out_of_period = [d for d in daterange(EXPORT_START, EXPORT_END) if not in_period(d)]
    eid, d = next((e, d) for d in out_of_period for e in pool if d not in used_dates[e])
    used_dates[eid].add(d)
    add(eid, at(d, minutes(7, 45)), None)                         # out of period
    cid = contractor_ids[0]
    d = free_dates(used_dates[cid], S.PERIOD_START, S.PERIOD_END, weekday_only=True, exclude_holidays=True)[0]
    used_dates[cid].add(d)
    add(cid, at(d, minutes(7, 50)), None)                         # contractor

    return boundary, missing_out_rows


def mangle_ids(rng, punches):
    variants = [
        lambda s: " " + s,
        lambda s: s + " ",
        lambda s: "  " + s + " ",
        lambda s: s.lower(),
        lambda s: s.lower() + " ",
        lambda s: " " + s.lower(),
    ]
    n = int(len(punches) * 0.04)
    for idx in rng.sample(range(len(punches)), n):
        punches[idx].employee_id = rng.choice(variants)(punches[idx].employee_id)
    return n


def add_duplicates(rng, punches, missing_out_rows):
    dup_idx = rng.sample(range(len(punches)), 25)
    dups = [punches[i].copy() for i in dup_idx]
    dups.append(missing_out_rows[0].copy())           # a duplicated *invalid* row
    dups.append(missing_out_rows[0].copy())           # ... three copies in total
    punches.extend(dups)
    return len(dups)


def sanity_check(punches, roster):
    """Guarantee that the export is unambiguous under the policy."""
    groups = defaultdict(list)
    for p in punches:
        key = (p.employee_id.strip().upper(), p.clock_in, p.clock_out, p.terminal)
        groups[key].append(p.row())
    for key, rows in groups.items():
        assert all(r == rows[0] for r in rows), f"inconsistent duplicate group {key}: {rows}"

    # same employee + same clock-in must never differ in clock-out/terminal
    seen = {}
    for key in groups:
        k2 = key[:2]
        assert k2 not in seen or seen[k2] == key[2:], f"conflicting rows for {k2}"
        seen[k2] = key[2:]

    # shifts of an employee never overlap in time
    per_emp = defaultdict(list)
    for key in groups:
        eid, ci, co, _ = key
        if co is not None and co > ci:
            per_emp[eid].append((ci, co))
    for eid, shifts in per_emp.items():
        shifts.sort()
        for (a_in, a_out), (b_in, b_out) in zip(shifts, shifts[1:]):
            assert b_in >= a_out, f"overlapping shifts for {eid}: {a_in}-{a_out} vs {b_in}-{b_out}"

    roster_ids = {e["Employee ID"] for e in roster}
    assert all(u not in roster_ids for u in UNKNOWN_IDS)


# --------------------------------------------------------------------------- #
# Writers
# --------------------------------------------------------------------------- #
HEADER_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")


def style_header(ws, ncols):
    for col in range(1, ncols + 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def write_roster(roster):
    wb = Workbook()
    ws = wb.active
    ws.title = S.ROSTER_SHEET
    ws.append(S.ROSTER_COLUMNS)
    for emp in roster:
        ws.append([emp[c] for c in S.ROSTER_COLUMNS])
    for r in range(2, len(roster) + 2):
        ws.cell(row=r, column=5).number_format = "0.00"
    style_header(ws, len(S.ROSTER_COLUMNS))
    for col, width in zip("ABCDEF", [14, 24, 20, 17, 12, 12]):
        ws.column_dimensions[col].width = width
    path = os.path.join(INITIAL_WORKSPACE, S.ROSTER_FILENAME)
    wb.save(path)
    return path


def write_attendance(punches):
    path = os.path.join(INITIAL_WORKSPACE, S.ATTENDANCE_FILENAME)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(S.ATTENDANCE_COLUMNS)
        for p in punches:
            w.writerow(p.row())
    return path


def write_holidays():
    path = os.path.join(INITIAL_WORKSPACE, S.HOLIDAYS_FILENAME)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(S.HOLIDAY_COLUMNS)
        for d, name in HOLIDAYS_2025:
            w.writerow([d, name])
    return path


def write_template():
    wb = Workbook()
    wb.remove(wb.active)
    widths = {
        S.SHEET_EMPLOYEE: [13, 24, 20, 17, 12, 8, 15, 14, 17, 15, 14, 21, 13, 14, 12],
        S.SHEET_DEPARTMENT: [20, 11, 14, 21, 13, 14, 12],
        S.SHEET_ALERTS: [13, 24, 20, 21, 14],
    }
    for sheet, columns in S.REPORT_SHEETS.items():
        ws = wb.create_sheet(sheet)
        ws.append(columns)
        style_header(ws, len(columns))
        for i, width in enumerate(widths[sheet]):
            ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].width = width
    path = os.path.join(INITIAL_WORKSPACE, S.TEMPLATE_FILENAME)
    wb.save(path)
    return path


# --------------------------------------------------------------------------- #
def main():
    rng = random.Random(SEED)
    os.makedirs(INITIAL_WORKSPACE, exist_ok=True)

    roster = build_roster(rng)
    profiles = assign_profiles(roster)
    punches, used_dates = generate_bulk(rng, roster, profiles)
    n_bulk = len(punches)
    boundary, missing_out_rows = add_engineered_rows(rng, roster, profiles, punches, used_dates)
    n_mangled = mangle_ids(rng, punches)
    n_dups = add_duplicates(rng, punches, missing_out_rows)
    rng.shuffle(punches)
    sanity_check(punches, roster)

    paths = [write_roster(roster), write_attendance(punches), write_holidays(), write_template()]

    eligible = [e for e in roster if e["Employment Type"] in S.ELIGIBLE_EMPLOYMENT_TYPES]
    print(f"seed                 : {SEED}")
    print(f"roster               : {len(roster)} employees, {len(eligible)} eligible")
    print(f"attendance rows      : {len(punches)} ({n_bulk} bulk, {n_mangled} mangled IDs, {n_dups} duplicates)")
    print("profiles             : " + ", ".join(
        f"{k}={sum(1 for v in profiles.values() if v == k)}"
        for k in ["day", "heavy", "night", "part", "split", "absent", "only_invalid"]))
    print("alert boundary       : " + ", ".join(f"{k} -> {v:.2f} h" for k, v in sorted(boundary.items())))
    for p in paths:
        print(f"wrote {os.path.relpath(p, os.path.dirname(HERE))}")


if __name__ == "__main__":
    main()
