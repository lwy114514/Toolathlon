#!/usr/bin/env python3
"""Reference solution for the overtime-payroll task.

Implements ``initial_workspace/overtime_policy.md`` literally and writes the
ground-truth workbook ``overtime_report_2025-08.xlsx`` next to this script.
The evaluator compares the agent's workbook against that file.

Money is computed with exact rational arithmetic (``fractions.Fraction``) so
the ground truth carries no floating-point drift.  ``generate_data.py``
guarantees that every pay figure is an exact number of cents (all hourly
rates have a cent value divisible by 8 and paid hours are quarter-hours), so
"round to two decimals" never has to break a tie.

Usage:
    python reference_solution.py
    python reference_solution.py --initial_workspace DIR --output FILE
"""

import argparse
import csv
import datetime as dt
import os
import sys
from collections import defaultdict
from fractions import Fraction

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import payroll_schema as S  # noqa: E402

DEFAULT_INITIAL_WORKSPACE = os.path.join(os.path.dirname(HERE), "initial_workspace")
DEFAULT_OUTPUT = os.path.join(HERE, S.OUTPUT_FILENAME)

MAX_RAW_MINUTES = 16 * 60
REGULAR_LIMIT = Fraction(S.REGULAR_HOURS_PER_SHIFT)
MULT_WEEKDAY_OT = Fraction(S.WEEKDAY_OT_MULTIPLIER)
MULT_WEEKEND = Fraction(S.WEEKEND_MULTIPLIER)
MULT_HOLIDAY = Fraction(S.HOLIDAY_MULTIPLIER)
ALERT_THRESHOLD = Fraction(S.OT_ALERT_THRESHOLD_HOURS)


# --------------------------------------------------------------------------- #
# Loading
# --------------------------------------------------------------------------- #
def load_roster(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[S.ROSTER_SHEET]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]
    if header != S.ROSTER_COLUMNS:
        raise ValueError(f"unexpected roster header: {header}")
    roster = []
    for row in rows:
        if row[0] is None:
            continue
        rec = dict(zip(header, row))
        rec["Employee ID"] = str(rec["Employee ID"]).strip().upper()
        rec["Hourly Rate"] = Fraction(round(float(rec["Hourly Rate"]) * 100), 100)
        roster.append(rec)
    wb.close()
    return roster


def load_holidays(path):
    with open(path, newline="", encoding="utf-8") as f:
        return {dt.date.fromisoformat(row["date"].strip()) for row in csv.DictReader(f)}


def parse_ts(value):
    value = (value or "").strip()
    return dt.datetime.strptime(value, "%Y-%m-%d %H:%M") if value else None


def load_punches(path):
    """Read the raw export, counting rows that are identical in every column once (policy 2.1)."""
    seen = set()
    punches = []
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        if [h.strip() for h in reader.fieldnames] != S.ATTENDANCE_COLUMNS:
            raise ValueError(f"unexpected attendance header: {reader.fieldnames}")
        for row in reader:
            key = tuple(row[c] for c in S.ATTENDANCE_COLUMNS)
            if key in seen:
                continue
            seen.add(key)
            punches.append({
                "employee_id": row["employee_id"].strip().upper(),   # policy 2.2
                "clock_in": parse_ts(row["clock_in"]),
                "clock_out": parse_ts(row["clock_out"]),
            })
    return punches


# --------------------------------------------------------------------------- #
# Policy
# --------------------------------------------------------------------------- #
def paid_hours(raw_minutes):
    """Policy 5: unpaid meal break, then round to the nearest quarter hour."""
    if raw_minutes > S.MEAL_BREAK_THRESHOLD_MINUTES:
        raw_minutes -= S.MEAL_BREAK_MINUTES
    # raw_minutes is an integer, so raw/15 can never sit exactly on x.5
    quarters = round(raw_minutes / S.ROUNDING_MINUTES)
    return Fraction(quarters * S.ROUNDING_MINUTES, 60)


def day_type(date, holidays):
    """Policy 6: holiday beats weekend beats weekday."""
    if date in holidays:
        return "holiday"
    if date.weekday() >= 5:
        return "weekend"
    return "weekday"


def in_period(date):
    return S.PERIOD_START <= date <= S.PERIOD_END


def compute(roster, holidays, punches):
    eligible = [e for e in roster if e["Employment Type"] in S.ELIGIBLE_EMPLOYMENT_TYPES]
    by_id = {e["Employee ID"]: e for e in eligible}

    stats = {
        eid: {
            "shifts": 0,
            "invalid": 0,
            "weekday_by_date": defaultdict(Fraction),
            "weekend": Fraction(0),
            "holiday": Fraction(0),
        }
        for eid in by_id
    }

    for p in punches:
        eid = p["employee_id"]
        if eid not in by_id:                       # policy 1: ineligible or not on roster
            continue
        ci, co = p["clock_in"], p["clock_out"]
        if not in_period(ci.date()):               # policy 3: date of clock-in decides
            continue
        st = stats[eid]
        if co is None or co <= ci:                 # policy 4
            st["invalid"] += 1
            continue
        raw = int((co - ci).total_seconds() // 60)
        if raw > MAX_RAW_MINUTES:                  # policy 4: badge left open
            st["invalid"] += 1
            continue
        st["shifts"] += 1
        hours = paid_hours(raw)
        kind = day_type(ci.date(), holidays)
        if kind == "weekday":
            st["weekday_by_date"][ci.date()] += hours   # policy 6: OT assessed per calendar day
        elif kind == "weekend":
            st["weekend"] += hours
        else:
            st["holiday"] += hours

    employee_rows = []
    for eid in sorted(by_id):
        e = by_id[eid]
        st = stats[eid]
        regular = sum((min(v, REGULAR_LIMIT) for v in st["weekday_by_date"].values()), Fraction(0))
        weekday_ot = sum((max(v - REGULAR_LIMIT, Fraction(0)) for v in st["weekday_by_date"].values()), Fraction(0))
        weekend, holiday = st["weekend"], st["holiday"]
        total_ot = weekday_ot + weekend + holiday
        rate = e["Hourly Rate"]
        regular_pay = rate * regular
        overtime_pay = rate * (MULT_WEEKDAY_OT * weekday_ot + MULT_WEEKEND * weekend + MULT_HOLIDAY * holiday)
        employee_rows.append({
            "Employee ID": eid,
            "Name": e["Name"],
            "Department": e["Department"],
            "Employment Type": e["Employment Type"],
            "Hourly Rate": rate,
            "Shifts": st["shifts"],
            "Invalid Punches": st["invalid"],
            "Regular Hours": regular,
            "Weekday OT Hours": weekday_ot,
            "Weekend Hours": weekend,
            "Holiday Hours": holiday,
            "Total Overtime Hours": total_ot,
            "Regular Pay": regular_pay,
            "Overtime Pay": overtime_pay,
            "Gross Pay": regular_pay + overtime_pay,
        })

    departments = defaultdict(list)
    for row in employee_rows:
        departments[row["Department"]].append(row)
    department_rows = []
    for dept in sorted(departments):
        rows = departments[dept]
        department_rows.append({
            "Department": dept,
            "Headcount": len(rows),
            "Regular Hours": sum((r["Regular Hours"] for r in rows), Fraction(0)),
            "Total Overtime Hours": sum((r["Total Overtime Hours"] for r in rows), Fraction(0)),
            "Regular Pay": sum((r["Regular Pay"] for r in rows), Fraction(0)),
            "Overtime Pay": sum((r["Overtime Pay"] for r in rows), Fraction(0)),
            "Gross Pay": sum((r["Gross Pay"] for r in rows), Fraction(0)),
        })

    alert_rows = [
        {c: r[c] for c in S.ALERT_COLUMNS}
        for r in employee_rows
        if r["Total Overtime Hours"] > ALERT_THRESHOLD
    ]
    alert_rows.sort(key=lambda r: (-r["Total Overtime Hours"], r["Employee ID"]))

    return {
        S.SHEET_EMPLOYEE: employee_rows,
        S.SHEET_DEPARTMENT: department_rows,
        S.SHEET_ALERTS: alert_rows,
    }


# --------------------------------------------------------------------------- #
# Output
# --------------------------------------------------------------------------- #
def to_cell(value):
    if isinstance(value, Fraction):
        return round(float(value), 2)
    return value


def write_report(report, template_path, output_path):
    wb = load_workbook(template_path)
    for sheet, columns in S.REPORT_SHEETS.items():
        ws = wb[sheet]
        header = [str(c.value).strip() for c in ws[1][: len(columns)]]
        if header != columns:
            raise ValueError(f"template sheet {sheet!r} header mismatch: {header}")
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for row in report[sheet]:
            ws.append([to_cell(row[c]) for c in columns])
        for r in range(2, ws.max_row + 1):
            for c_idx, col in enumerate(columns, start=1):
                if isinstance(report[sheet][r - 2][col], Fraction):
                    ws.cell(row=r, column=c_idx).number_format = "0.00"
    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="Build the ground-truth overtime report")
    parser.add_argument("--initial_workspace", default=DEFAULT_INITIAL_WORKSPACE)
    parser.add_argument("--output", default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    ws_dir = args.initial_workspace
    roster = load_roster(os.path.join(ws_dir, S.ROSTER_FILENAME))
    holidays = load_holidays(os.path.join(ws_dir, S.HOLIDAYS_FILENAME))
    punches = load_punches(os.path.join(ws_dir, S.ATTENDANCE_FILENAME))

    report = compute(roster, holidays, punches)
    write_report(report, os.path.join(ws_dir, S.TEMPLATE_FILENAME), args.output)

    emp = report[S.SHEET_EMPLOYEE]
    print(f"roster: {len(roster)} employees, {len(emp)} eligible")
    print(f"punches after de-duplication: {len(punches)}")
    print(f"valid shifts: {sum(r['Shifts'] for r in emp)}, invalid punches: {sum(r['Invalid Punches'] for r in emp)}")
    print(f"total overtime hours: {float(sum(r['Total Overtime Hours'] for r in emp)):.2f}")
    print(f"gross payroll: {float(sum(r['Gross Pay'] for r in emp)):,.2f}")
    print(f"departments: {len(report[S.SHEET_DEPARTMENT])}, OT alerts: {len(report[S.SHEET_ALERTS])}")
    for r in report[S.SHEET_ALERTS]:
        print(f"  alert {r['Employee ID']} {r['Name']:<22} {float(r['Total Overtime Hours']):6.2f} h")
    print(f"wrote {args.output}")


if __name__ == "__main__":
    main()
