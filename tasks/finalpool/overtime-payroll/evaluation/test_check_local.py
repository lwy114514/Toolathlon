#!/usr/bin/env python3
"""Self-test for evaluation/check_local.py.

Builds throw-away agent workspaces from the ground-truth workbook, mutates
them in the ways a real agent tends to get wrong, and asserts that the checker
accepts exactly the correct variants.

Run from the repository root or the task directory:
    python tasks/finalpool/overtime-payroll/evaluation/test_check_local.py
"""

import os
import shutil
import sys
import tempfile

import pandas as pd
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from check_local import OUTPUT_FILENAME, check_local  # noqa: E402

TASK_DIR = os.path.dirname(HERE)
GT_DIR = os.path.join(TASK_DIR, "groundtruth_workspace")
GT_FILE = os.path.join(GT_DIR, OUTPUT_FILENAME)

results = []


def case(name, expect_pass, builder):
    tmp = tempfile.mkdtemp(prefix="otp_")
    try:
        builder(tmp)
        ok, msg = check_local(tmp, GT_DIR)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    status = "ok " if ok == expect_pass else "BAD"
    results.append(ok == expect_pass)
    print(f"[{status}] {name:<48} -> {'pass' if ok else 'fail'}: {msg[:110]}")


def copy_gt(tmp):
    shutil.copy(GT_FILE, os.path.join(tmp, OUTPUT_FILENAME))
    return os.path.join(tmp, OUTPUT_FILENAME)


def mutate(fn):
    """Return a builder that copies the GT and applies fn(workbook)."""
    def build(tmp):
        path = copy_gt(tmp)
        wb = load_workbook(path)
        fn(wb)
        wb.save(path)
    return build


# ----------------------------------------------------------------- positives
case("exact ground truth", True, copy_gt)


def via_pandas(tmp):
    sheets = pd.read_excel(GT_FILE, sheet_name=None)
    with pd.ExcelWriter(os.path.join(tmp, OUTPUT_FILENAME), engine="openpyxl") as w:
        for name, df in sheets.items():
            df.to_excel(w, sheet_name=name, index=False)


case("rewritten with pandas (no styling)", True, via_pandas)


def numbers_as_text(wb):
    ws = wb["Employee Summary"]
    for r in range(2, ws.max_row + 1):
        for c in range(5, ws.max_column + 1):
            ws.cell(r, c).value = f"{ws.cell(r, c).value}"


case("numeric cells stored as text", True, mutate(numbers_as_text))


def within_tolerance(wb):
    ws = wb["Employee Summary"]
    ws.cell(2, 15).value = float(ws.cell(2, 15).value) + 0.004


case("value off by 0.004 (inside tolerance)", True, mutate(within_tolerance))


def extra_sheet(wb):
    ws = wb.create_sheet("Scratch")
    ws.append(["notes"])


case("extra scratch sheet", True, mutate(extra_sheet))


def integer_as_float(wb):
    ws = wb["Employee Summary"]
    ws.cell(2, 6).value = float(ws.cell(2, 6).value)


case("Shifts written as 18.0 instead of 18", True, mutate(integer_as_float))


def whitespace_in_text(wb):
    ws = wb["Employee Summary"]
    ws.cell(2, 2).value = "  " + ws.cell(2, 2).value + " "
    ws.cell(1, 1).value = "Employee ID "


case("stray whitespace in text cells/header", True, mutate(whitespace_in_text))


# ----------------------------------------------------------------- negatives
case("output file missing", False, lambda tmp: None)


def wrong_name(tmp):
    shutil.copy(GT_FILE, os.path.join(tmp, "overtime_report.xlsx"))


case("wrong file name", False, wrong_name)


def template_only(tmp):
    shutil.copy(os.path.join(TASK_DIR, "initial_workspace", "overtime_report_template.xlsx"),
                os.path.join(tmp, OUTPUT_FILENAME))


case("template saved without data", False, template_only)


def drop_sheet(wb):
    wb.remove(wb["OT Alerts"])


case("OT Alerts sheet missing", False, mutate(drop_sheet))


def rename_header(wb):
    wb["Department Summary"].cell(1, 2).value = "Head Count"


case("header renamed", False, mutate(rename_header))


def swap_columns(wb):
    ws = wb["Employee Summary"]
    for r in range(1, ws.max_row + 1):
        a, b = ws.cell(r, 13).value, ws.cell(r, 14).value
        ws.cell(r, 13).value, ws.cell(r, 14).value = b, a


case("two columns swapped", False, mutate(swap_columns))


def off_by_cent(wb):
    ws = wb["Employee Summary"]
    ws.cell(3, 15).value = float(ws.cell(3, 15).value) + 0.01


case("Gross Pay off by one cent", False, mutate(off_by_cent))


def wrong_shift_count(wb):
    ws = wb["Employee Summary"]
    ws.cell(4, 6).value = int(ws.cell(4, 6).value) + 1


case("Shifts off by one", False, mutate(wrong_shift_count))


def sort_by_name(wb):
    ws = wb["Employee Summary"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    rows.sort(key=lambda r: r[1])
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j).value = v


case("Employee Summary sorted by Name", False, mutate(sort_by_name))


def alerts_ascending(wb):
    ws = wb["OT Alerts"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    rows.reverse()
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j).value = v


case("OT Alerts sorted ascending", False, mutate(alerts_ascending))


def missing_row(wb):
    wb["Employee Summary"].delete_rows(10)


case("an employee row dropped", False, mutate(missing_row))


def extra_alert(wb):
    ws = wb["OT Alerts"]
    ws.append(["EMP-1017", "Someone", "Assembly", 30.0, 1000.0])


case("employee at exactly 30.00 h added to alerts", False, mutate(extra_alert))


def formula_cell(wb):
    ws = wb["Employee Summary"]
    ws.cell(2, 15).value = "=M2+N2"


case("Gross Pay written as a formula", False, mutate(formula_cell))


def empty_cell(wb):
    wb["Department Summary"].cell(2, 7).value = None


case("a numeric cell left empty", False, mutate(empty_cell))


def text_in_number(wb):
    wb["Employee Summary"].cell(5, 8).value = "n/a"


case("text where a number is expected", False, mutate(text_in_number))


def wrong_department_name(wb):
    wb["Department Summary"].cell(2, 1).value = "Admin"


case("department name altered", False, mutate(wrong_department_name))

print()
good = sum(results)
print(f"{good}/{len(results)} checker cases behaved as expected")
sys.exit(0 if good == len(results) else 1)
